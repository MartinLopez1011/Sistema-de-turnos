"""
Generador del Documento de Requerimientos y Alcance del Sistema de Turnos en formato Word (.docx).
Diseñado para ser simple, claro y entendible para usuarios sin conocimientos informáticos,
con una definición exhaustiva y transparente del alcance del proyecto.
"""

import os
import sys
from datetime import date

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

# Asegurar salida de consola UTF-8 en Windows
if sys.platform == "win32":
    try:
        sys.stdout.reconfigure(encoding="utf-8")
        sys.stderr.reconfigure(encoding="utf-8")
    except Exception:
        pass

ROOT = os.path.dirname(os.path.abspath(__file__))
OUTPUT = os.path.join(ROOT, "Requerimientos_del_Sistema_Turnos.docx")


def set_cell_shading(cell, fill_hex):
    """Aplica color de fondo a una celda de tabla."""
    properties = cell._tc.get_or_add_tcPr()
    shading = properties.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        properties.append(shading)
    shading.set(qn("w:fill"), fill_hex)


def set_cell_text(cell, text, bold=False, color="1F2937", size=9.5, align=WD_ALIGN_PARAGRAPH.LEFT):
    """Configura el texto, formato, fuente y alineación de una celda."""
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.alignment = align
    paragraph.paragraph_format.space_after = Pt(2.5)
    paragraph.paragraph_format.space_before = Pt(2.5)
    paragraph.paragraph_format.line_spacing = 1.08
    run = paragraph.add_run(str(text))
    run.bold = bold
    run.font.name = "Segoe UI"
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def set_table_borders(table, color="CBD5E1", size="6"):
    """Configura bordes delgados y elegantes en la tabla."""
    properties = table._tbl.tblPr
    borders = properties.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        properties.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = "w:" + edge
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def add_custom_table(document, headers, rows, widths=None):
    """Crea una tabla con encabezado institucional estilizado y alternancia de colores."""
    table = document.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    set_table_borders(table)

    # Fila de encabezado
    for index, header in enumerate(headers):
        align = WD_ALIGN_PARAGRAPH.CENTER if index == 0 else WD_ALIGN_PARAGRAPH.LEFT
        set_cell_text(table.rows[0].cells[index], header, bold=True, color="FFFFFF", size=9.5, align=align)
        set_cell_shading(table.rows[0].cells[index], "0F766E")

    # Filas de datos
    for r_idx, row in enumerate(rows):
        cells = table.add_row().cells
        bg_color = "F0FDFA" if r_idx % 2 == 1 else "FFFFFF"
        for index, value in enumerate(row):
            align = WD_ALIGN_PARAGRAPH.CENTER if index == 0 and len(str(value)) <= 8 else WD_ALIGN_PARAGRAPH.LEFT
            bold = True if index == 0 and len(str(value)) <= 8 else False
            set_cell_text(cells[index], value, bold=bold, size=9.2, align=align)
            set_cell_shading(cells[index], bg_color)

    if widths:
        for row in table.rows:
            for index, width in enumerate(widths):
                row.cells[index].width = Inches(width)

    p_after = document.add_paragraph()
    p_after.paragraph_format.space_after = Pt(4)
    return table


def add_callout_box(document, title, text, box_type="info"):
    """Crea un cuadro de aviso destacado y fácil de leer."""
    table = document.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    cell = table.rows[0].cells[0]
    cell.width = Inches(7.0)

    # Colores según tipo
    cfg = {
        "info": ("F0FDF4", "10B981", "065F46"),   # Verde menta
        "warn": ("FFFBEB", "F59E0B", "92400E"),   # Amarillo/Naranja
        "alert": ("FEF2F2", "EF4444", "991B1B"),  # Rojo suave
    }
    bg_color, border_color, text_color = cfg.get(box_type, cfg["info"])

    set_cell_shading(cell, bg_color)

    # Configurar borde izquierdo grueso y sin bordes en los otros lados
    properties = cell._tc.get_or_add_tcPr()
    borders = properties.find(qn("w:tcBorders"))
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        properties.append(borders)
    
    for edge in ("top", "bottom", "right"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "none")
        borders.append(el)
    
    left = OxmlElement("w:left")
    left.set(qn("w:val"), "single")
    left.set(qn("w:sz"), "24") # Grueso
    left.set(qn("w:color"), border_color)
    borders.append(left)

    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(2)
    run_title = p.add_run(f"📌 {title}\n")
    run_title.bold = True
    run_title.font.name = "Segoe UI"
    run_title.font.size = Pt(10)
    run_title.font.color.rgb = RGBColor.from_string(text_color)

    run_body = p.add_run(text)
    run_body.font.name = "Segoe UI"
    run_body.font.size = Pt(9.5)
    run_body.font.color.rgb = RGBColor(31, 41, 55)

    p_after = document.add_paragraph()
    p_after.paragraph_format.space_after = Pt(4)


def add_heading(document, text, level=1):
    paragraph = document.add_heading(text, level=level)
    paragraph.paragraph_format.space_before = Pt(14 if level == 1 else 10)
    paragraph.paragraph_format.space_after = Pt(4)
    return paragraph


def add_body(document, text, bold_lead=None):
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(6)
    paragraph.paragraph_format.line_spacing = 1.15
    if bold_lead and text.startswith(bold_lead):
        r_lead = paragraph.add_run(bold_lead)
        r_lead.bold = True
        r_lead.font.name = "Segoe UI"
        r_lead.font.size = Pt(10)
        r_lead.font.color.rgb = RGBColor(15, 23, 42)
        r_body = paragraph.add_run(text[len(bold_lead):])
        r_body.font.name = "Segoe UI"
        r_body.font.size = Pt(10)
        r_body.font.color.rgb = RGBColor(51, 65, 85)
    else:
        r = paragraph.add_run(text)
        r.font.name = "Segoe UI"
        r.font.size = Pt(10)
        r.font.color.rgb = RGBColor(51, 65, 85)
    return paragraph


def add_bullets(document, items):
    for item in items:
        paragraph = document.add_paragraph(style="List Bullet")
        paragraph.paragraph_format.space_after = Pt(3)
        paragraph.paragraph_format.line_spacing = 1.12
        if ":" in item:
            parts = item.split(":", 1)
            r1 = paragraph.add_run(parts[0] + ":")
            r1.bold = True
            r1.font.name = "Segoe UI"
            r1.font.size = Pt(9.8)
            r1.font.color.rgb = RGBColor(15, 23, 42)
            r2 = paragraph.add_run(parts[1])
            r2.font.name = "Segoe UI"
            r2.font.size = Pt(9.8)
            r2.font.color.rgb = RGBColor(51, 65, 85)
        else:
            r = paragraph.add_run(item)
            r.font.name = "Segoe UI"
            r.font.size = Pt(9.8)
            r.font.color.rgb = RGBColor(51, 65, 85)


def configure_document_styles(document):
    styles = document.styles
    normal = styles["Normal"]
    normal.font.name = "Segoe UI"
    normal.font.size = Pt(10)
    normal.font.color.rgb = RGBColor(31, 41, 55)

    for name, size, color in (
        ("Title", 24, "0F766E"),
        ("Heading 1", 14, "0F766E"),
        ("Heading 2", 11.5, "115E59"),
        ("Heading 3", 10.5, "1E293B")
    ):
        style = styles[name]
        style.font.name = "Segoe UI"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)


def add_header_footer(section):
    header = section.header.paragraphs[0]
    header.text = "Sistema de Gestión de Turnos  |  Requerimientos y Alcance del Sistema"
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    for run in header.runs:
        run.font.name = "Segoe UI"
        run.font.size = Pt(8.5)
        run.font.color.rgb = RGBColor(107, 114, 128)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in footer.runs:
        run.font.name = "Segoe UI"
        run.font.size = Pt(8.5)
        run.font.color.rgb = RGBColor(107, 114, 128)
    
    # Campo de página Word
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)


def build_word_document():
    document = Document()
    configure_document_styles(document)

    section = document.sections[0]
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.75)
    section.right_margin = Inches(0.75)
    add_header_footer(section)

    document.core_properties.title = "Requerimientos y Alcance - Sistema de Gestión de Turnos"
    document.core_properties.subject = "Documento de Requerimientos y Alcance en Lenguaje Simple"
    document.core_properties.author = "Sistema de Gestión de Turnos"

    # ── TÍTULO PRINCIPAL ──────────────────────────────────────────────────────────
    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_before = Pt(16)
    title.paragraph_format.space_after = Pt(2)
    run_t = title.add_run("Documento de Requerimientos y Alcance")
    run_t.bold = True
    run_t.font.name = "Segoe UI"
    run_t.font.size = Pt(22)
    run_t.font.color.rgb = RGBColor(15, 118, 110)

    sub = document.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub.paragraph_format.space_after = Pt(16)
    run_sub = sub.add_run("Sistema de Planificación de Turnos de Guardia Semanal — Guía Clara y Sin Tecnicismos")
    run_sub.font.name = "Segoe UI"
    run_sub.font.size = Pt(11.5)
    run_sub.font.color.rgb = RGBColor(75, 85, 99)

    # Ficha Técnica
    ficha_data = [
        ("Nombre del Sistema", "Sistema de Gestión de Turnos (TurnosApp)"),
        ("Tipo de Aplicación", "Programa de escritorio para computador (Windows)"),
        ("Destinatarios", "Jefaturas de Servicio, Coordinadores de Guardia y Funcionarios"),
        ("Objetivo Principal", "Calcular y ordenar las guardias semanales de forma automática, justa y transparente"),
        ("Fecha de Documento", date.today().strftime("%d/%m/%Y")),
    ]
    ficha_table = document.add_table(rows=len(ficha_data), cols=2)
    ficha_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    ficha_table.style = "Table Grid"
    set_table_borders(ficha_table, color="CBD5E1")
    for idx, (k, v) in enumerate(ficha_data):
        set_cell_text(ficha_table.rows[idx].cells[0], k, bold=True, color="FFFFFF", size=9.5)
        set_cell_shading(ficha_table.rows[idx].cells[0], "0F766E")
        set_cell_text(ficha_table.rows[idx].cells[1], v, size=9.5)
        set_cell_shading(ficha_table.rows[idx].cells[1], "F8FAFC")
        ficha_table.rows[idx].cells[0].width = Inches(2.0)
        ficha_table.rows[idx].cells[1].width = Inches(5.0)

    p_sp = document.add_paragraph()
    p_sp.paragraph_format.space_after = Pt(6)

    # ── 1. ¿QUÉ ES EL SISTEMA Y QUÉ PROBLEMA RESUELVE? ───────────────────────────
    add_heading(document, "1. ¿Qué es este sistema y qué necesidad resuelve?", level=1)
    add_body(
        document,
        "Antes de este sistema, los turnos de guardia se calculaban a mano en hojas de papel o en planillas de Excel aisladas. Esto generaba problemas frecuentes: personas a las que se les asignaba turno mientras estaban de vacaciones o con licencia médica, funcionarios que repetían guardias pesadas (como Navidad o Año Nuevo) dos años seguidos, y desacuerdos sobre a quién le correspondía realmente el turno."
    )
    add_body(
        document,
        "El Sistema de Gestión de Turnos es un programa para el computador que automatiza completamente este proceso. Actúa como un árbitro neutral: lleva la lista ordenada de todo el equipo, calcula a quién le corresponde cada semana, salta de manera automática a quien tenga un permiso justificado y genera la planilla oficial de Excel lista para imprimir y firmar."
    )

    add_callout_box(
        document,
        "Principio de Equidad y Transparencia",
        "El programa garantiza que la carga de trabajo se reparta de forma equitativa entre todos los funcionarios. Nadie trabaja turnos seguidos injustamente y nadie queda exento sin justificación.",
        "info"
    )

    # ── 2. DEFINICIÓN DEL ALCANCE (QUÉ INCLUYE Y QUÉ NO) ────────────────────────
    add_heading(document, "2. Definición Detallada del Alcance", level=1)
    add_body(
        document,
        "El alcance define con absoluta claridad la frontera del sistema: exactamente qué responsabilidades asume el programa y cuáles quedan fuera de su función."
    )

    alcance_items = [
        (
            "Lo que SÍ hace el sistema\n(Dentro del Alcance)",
            "• Asignación Semanal de Guardias: Calcula quién cubre la guardia de lunes a domingo para cada semana del mes.\n"
            "• Rotación Circular Equitativa: Sigue el orden estricto de la lista de personal para que todos hagan turno por igual.\n"
            "• Gestión de Permisos y Ausencias: Registra Días Administrativos (DA), Vacaciones (FL), Licencias Médicas (LIC) u Otros Permisos (OTR) e inhabilita al funcionario durante su ausencia.\n"
            "• Sistema de Compensación (Lista de Pendientes): Si un funcionario no pudo hacer su guardia por permiso, el sistema no le 'perdona' el turno: lo anota en una lista de espera prioritaria para que recupere la guardia en cuanto vuelva a estar disponible.\n"
            "• Cuidado de Fiestas de Diciembre: Revisa automáticamente si la persona ya trabajó en Navidad o Año Nuevo el año anterior para evitar que repita el mismo feriado.\n"
            "• Cambios Manuales (Permutas de palabra): Permite al coordinador cambiar a mano una guardia solicitando obligatoriamente el motivo del cambio.\n"
            "• Previsualización no destructiva: Permite mirar y jugar con los meses futuros sin alterar el orden guardado.\n"
            "• Cierre Definitivo del Mes: Guarda el historial consolidado y hace girar la lista para el mes siguiente.\n"
            "• Planilla Oficial en Excel (.xlsx): Exporta el calendario mensual completo con colores, leyenda y totales de turnos por persona, listo para imprimir.\n"
            "• Administración del Personal: Permite agregar nuevos compañeros, corregir nombres, actualizar correos, cambiar el orden de la fila y retirar a quienes se trasladan sin borrar su historial pasado.\n"
            "• Aviso por Correo Electrónico: Si se hace un cambio manual a última hora, al cerrar el mes envía automáticamente un correo a todo el equipo informando el cambio.\n"
            "• Copias de Seguridad Automáticas: Guarda respaldos continuos en una carpeta /backups para que nunca se pierda la información si se corta la luz."
        ),
        (
            "Lo que NO hace el sistema\n(Fuera del Alcance)",
            "• No es un reloj control de asistencia diaria: No registra horas de llegada ni de salida del personal.\n"
            "• No aprueba permisos administrativos ni vacaciones: El coordinador debe ingresar los permisos al sistema una vez que hayan sido autorizados por la jefatura.\n"
            "• No requiere internet para el cálculo: El programa funciona 100% de forma local y desconectada en el computador (solo usa internet si se envía el correo por cambio manual).\n"
            "• No es un sistema multiusuario en red con contraseñas: Está pensado para ser operado por el encargado o coordinador de la unidad en su equipo de trabajo.\n"
            "• No envía mensajes de WhatsApp ni SMS: La comunicación oficial de cambios manuales se realiza exclusivamente por correo electrónico vía Gmail institucional.\n"
            "• No modifica meses pasados por error: Un mes cerrado queda protegido como historial inmutable."
        )
    ]
    add_custom_table(document, ["Clasificación", "Detalle Exhaustivo del Alcance"], alcance_items, widths=[2.2, 4.8])

    # ── 3. QUIÉNES USAN EL SISTEMA (ACTORES) ─────────────────────────────────────
    add_heading(document, "3. Usuarios del Sistema (Roles)", level=1)
    add_bullets(document, [
        "El Coordinador o Encargado de Turnos: Es la persona que maneja el programa. Elige el mes, ingresa los permisos conocidos, revisa que las semanas estén correctas, saca la planilla en Excel para firmar y guarda el mes definitivo.",
        "La Jefatura de la Unidad: Recibe la planilla Excel mensual oficial para su visado, firma y publicación.",
        "Los Funcionarios del Equipo: Son los beneficiarios de la rotación justa. Pueden consultar sus turnos en el calendario Excel y reciben un correo de aviso si su guardia fue modificada de mutuo acuerdo."
    ])

    # ── 4. REQUERIMIENTOS FUNCIONALES (LO QUE HACE LA PANTALLA) ─────────────────
    add_heading(document, "4. Requerimientos Funcionales (Lo que hace cada botón y pantalla)", level=1)
    add_body(
        document,
        "A continuación se describen las capacidades que el sistema ofrece al usuario en su uso cotidiano, explicadas de forma sencilla y directa:"
    )

    rf_tabla = [
        ("RF-01", "Elegir Mes y Año", "El usuario puede seleccionar libremente cualquier mes y año en pantalla. El sistema carga de inmediato las semanas correspondientes a ese período."),
        ("RF-02", "Rotación Automática", "El sistema reparte las semanas de guardia en orden secuencial entre todos los funcionarios de la lista, asegurando que a todos les toque la misma cantidad."),
        ("RF-03", "Ingresar Ausencias", "El usuario puede ingresar días o períodos de ausencia para cualquier funcionario escribiendo los números (ej. '1-5' o '12, 19') y seleccionando si es Día Administrativo (DA), Vacaciones (FL), Licencia Médica (LIC) u Otro (OTR)."),
        ("RF-04", "Salto y Devolución Justa", "Si alguien tiene permiso, el programa no le da turno esa semana, le asigna al siguiente compañero libre y guarda a la persona ausente en una 'lista de pendientes' para devolverle la guardia apenas regrese."),
        ("RF-05", "Cambio Manual de Guardia", "Si dos personas cambiaron turno entre ellas, el usuario puede pulsar 'Cambiar' en la semana, seleccionar al reemplazante y escribir obligatoriamente el motivo del cambio (queda marcado con la etiqueta verde MANUAL)."),
        ("RF-06", "Botón para Volver a Auto", "Si el usuario se equivoca al cambiar una guardia a mano, puede pulsar el botón '↺ Auto' para que el sistema vuelva a calcular la persona original que le tocaba por lista."),
        ("RF-07", "Protección Fiestas de Diciembre", "En diciembre, el sistema consulta los feriados nacionales de Chile para evitar que un funcionario repita Navidad o Año Nuevo si ya le tocó hacer guardia en esa misma fiesta el año anterior."),
        ("RF-08", "Calendario Mensual a Color", "En la pestaña 'Ver Turnos del Mes', el usuario puede ver la cuadrícula completa del mes: días de turno en rojo, ausencias en sus colores y fines de semana sombreados."),
        ("RF-09", "Exportar a Excel Oficial", "El usuario puede pulsar 'Exportar Excel' para generar un archivo .xlsx con título institucional, nombres completos, días coloreados, tabla de totales por persona y leyenda de colores."),
        ("RF-10", "Guardar Mes Definitivo", "Al pulsar 'Guardar mes' en Planificación, el programa guarda el período en el historial, avanza la cola de rotación para el siguiente mes y deja todo listo para continuar."),
        ("RF-11", "Administrar la Lista del Equipo", "En la pestaña Ajustes, el usuario puede agregar nuevas personas con su correo, corregir nombres, reordenar la fila con flechas (⬆/⬇) y dar de baja a quienes dejen la unidad."),
        ("RF-12", "Avisos por Correo", "Al guardar un mes que contenga cambios manuales, el sistema envía automáticamente un correo formal a todos los funcionarios detallando la semana cambiada y el motivo."),
    ]
    add_custom_table(document, ["ID", "Funcionalidad", "¿Qué le permite hacer al usuario en la práctica?"], rf_tabla, widths=[0.8, 1.9, 4.3])

    # ── 5. REQUERIMIENTOS DE OPERACIÓN Y FUNCIONAMIENTO ─────────────────────────
    add_heading(document, "5. Requerimientos de Operación (¿Qué necesita para funcionar?)", level=1)
    
    rnf_tabla = [
        ("Velocidad Inmediata", "El cálculo de las semanas y la vista previa del calendario toman menos de 1 segundo. La creación de la hoja de Excel toma menos de 2 segundos."),
        ("Diseño Amigable (Modo Oscuro)", "La pantalla tiene fondo oscuro descansado para la vista, letras grandes, avatares con las iniciales de cada persona y avisos en verde o rojo según corresponda."),
        ("Sin Instalaciones Complicadas", "Viene como un único archivo ejecutable ('Sistema de Turnos.exe'). No requiere instalar Python ni programas técnicos adicionales en el computador."),
        ("Trabajo Desconectado (Offline)", "El programa no depende de internet para armar los turnos ni para sacar las planillas de Excel."),
        ("Seguridad y Respaldo", "El programa guarda su información de forma segura en 'config.json' y crea copias de respaldo automáticas fechadas en la carpeta /backups cada vez que se guarda un mes."),
    ]
    add_custom_table(document, ["Aspecto Clave", "Condición Práctica"], rnf_tabla, widths=[2.0, 5.0])

    # ── 6. REGLAS DE NEGOCIO EN PALABRAS SIMPLES ─────────────────────────────────
    add_heading(document, "6. Reglas de Negocio en Palabras Simples", level=1)
    add_body(
        document,
        "Estas son las 3 reglas básicas con las que piensa el programa para que todo sea justo:"
    )

    add_bullets(document, [
        "Regla 1 (El orden para elegir quién hace guardia): Para cada semana, el programa revisa en este orden: 1° Si el coordinador fijó a alguien a mano (asignación manual). 2° Si hay semanas fijas de inicio. 3° Si el mes ya estaba guardado en el historial. 4° Si alguien debe turno por haber estado de permiso (lista de pendientes). 5° Si nadie debe turno, le toca al siguiente de la lista en orden circular.",
        "Regla 2 (Exportar Excel NO es lo mismo que Guardar Mes): Exportar Excel solo crea la hoja para mirar, imprimir o mandar a borrador (puedes exportar cuantas veces quieras y la cuenta no se mueve). Guardar Mes es la firma final que anota el mes en el historial y hace girar la lista para el mes siguiente.",
        "Regla 3 (Nadie hace guardia estando de permiso): Si una persona tiene anotado un Día Administrativo (DA), Feriado Legal (FL), Licencia Médica (LIC) u Otro (OTR) en los días de una semana, el sistema jamás le asignará guardia en esa semana."
    ])

    # ── 7. VALIDACIÓN Y APROBACIÓN DEL ALCANCE ───────────────────────────────────
    add_heading(document, "7. Ficha de Aprobación y Conformidad del Alcance", level=1)
    add_body(
        document,
        "La firma del presente documento acredita que los requerimientos y el alcance aquí descritos representan fielmente las necesidades operativas de la unidad para la gestión de turnos de guardia:"
    )

    app_table = document.add_table(rows=3, cols=3)
    app_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    app_table.style = "Table Grid"
    set_table_borders(app_table, color="CBD5E1")

    headers_app = ["Rol / Cargo", "Nombre del Responsable", "Firma y Fecha"]
    for i, h in enumerate(headers_app):
        set_cell_text(app_table.rows[0].cells[i], h, bold=True, color="FFFFFF", size=9.5, align=WD_ALIGN_PARAGRAPH.CENTER)
        set_cell_shading(app_table.rows[0].cells[i], "0F766E")

    roles = [
        ("Coordinador de Turnos", "", ""),
        ("Jefatura de Unidad / Servicio", "", ""),
    ]
    for row_idx, (rol, nom, f) in enumerate(roles, start=1):
        cells = app_table.rows[row_idx].cells
        set_cell_text(cells[0], rol, bold=True, size=9.5, align=WD_ALIGN_PARAGRAPH.CENTER)
        set_cell_shading(cells[0], "F8FAFC")
        set_cell_text(cells[1], nom, size=9.5)
        set_cell_text(cells[2], f, size=9.5)

    for row in app_table.rows:
        row.cells[0].width = Inches(2.2)
        row.cells[1].width = Inches(2.8)
        row.cells[2].width = Inches(2.0)

    document.save(OUTPUT)
    print(f"[OK] Documento de requerimientos generado exitosamente en: {OUTPUT}")


if __name__ == "__main__":
    build_word_document()
