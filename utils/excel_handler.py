import calendar
import datetime
from datetime import timedelta
import openpyxl
import os
from openpyxl.styles import PatternFill, Alignment, Border, Side, Font
from openpyxl.comments import Comment

class ExcelHandler:
    def __init__(self, output_path, personal):
        self.output_path = output_path
        self.personal = personal
        self.wb = None
        self.sheet = None

    def load_template(self):
        self.wb = openpyxl.Workbook()
        self.sheet = self.wb.active
        self.sheet.title = "Turnos"

        title_cell = self.sheet["A1"]
        title_cell.value = "PLANIFICACION DE TURNOS"
        title_cell.font = Font(bold=True, size=14)
        title_cell.alignment = Alignment(horizontal="center", vertical="center")

        header_fill = PatternFill(start_color="D9E2F3", end_color="D9E2F3", fill_type="solid")
        border = Border(
            left=Side(style="thin", color="808080"),
            right=Side(style="thin", color="808080"),
            top=Side(style="thin", color="808080"),
            bottom=Side(style="thin", color="808080"),
        )
        self.sheet["A2"] = "DIA"
        self.sheet["A3"] = "FUNCIONARIO"
        for row in (2, 3):
            self.sheet.cell(row=row, column=1).font = Font(bold=True)
            self.sheet.cell(row=row, column=1).fill = header_fill
            self.sheet.cell(row=row, column=1).alignment = Alignment(horizontal="center", vertical="center")
            self.sheet.cell(row=row, column=1).border = border

        for day in range(1, 32):
            cell = self.sheet.cell(row=2, column=day + 1, value=day)
            cell.font = Font(bold=True)
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = border

            cell3 = self.sheet.cell(row=3, column=day + 1)
            cell3.font = Font(bold=True)
            cell3.fill = header_fill
            cell3.alignment = Alignment(horizontal="center", vertical="center")
            cell3.border = border

        # Encabezados de columnas de resumen / totales
        summary_headers = [
            (33, "AG", "TURNOS", 10),
            (34, "AH", "DA", 6),
            (35, "AI", "FL", 6),
            (36, "AJ", "LIC", 6),
            (37, "AK", "OTR", 6),
        ]
        self.sheet.merge_cells("AG2:AK2")
        sum_title = self.sheet.cell(row=2, column=33, value="TOTALES")
        sum_title.font = Font(bold=True)
        sum_title.fill = header_fill
        sum_title.alignment = Alignment(horizontal="center", vertical="center")
        for col in range(33, 38):
            self.sheet.cell(row=2, column=col).border = border
            self.sheet.cell(row=2, column=col).fill = header_fill

        for col_idx, col_letter, label, col_w in summary_headers:
            cell_sub = self.sheet.cell(row=3, column=col_idx, value=label)
            cell_sub.font = Font(bold=True, size=9)
            cell_sub.fill = header_fill
            cell_sub.alignment = Alignment(horizontal="center", vertical="center")
            cell_sub.border = border
            self.sheet.column_dimensions[col_letter].width = col_w

        for row_idx, person in enumerate(self.personal, 4):
            name = person.get("nombre", "") if isinstance(person, dict) else str(person)
            cell = self.sheet.cell(row=row_idx, column=1, value=name)
            cell.border = border
            for column_idx in range(2, 38):
                self.sheet.cell(row=row_idx, column=column_idx).border = border

        # Ancho dinámico para la columna A (nombres de funcionarios)
        max_name_len = 0
        for p in self.personal:
            p_name = p.get("nombre", "") if isinstance(p, dict) else str(p)
            if len(p_name) > max_name_len:
                max_name_len = len(p_name)
        self.sheet.column_dimensions["A"].width = max(38, max_name_len + 4)

        for column_idx in range(2, 33):
            self.sheet.column_dimensions[openpyxl.utils.get_column_letter(column_idx)].width = 5
        self.sheet.freeze_panes = "B4"
        self.sheet.row_dimensions[1].height = 24

        # Configuración de impresión: apaisado y ajustado a 1 página de ancho
        self.sheet.page_setup.orientation = self.sheet.ORIENTATION_LANDSCAPE
        self.sheet.page_setup.paperSize = self.sheet.PAPERSIZE_LETTER
        self.sheet.sheet_properties.pageSetUpPr.fitToPage = True
        self.sheet.page_setup.fitToWidth = 1
        self.sheet.page_setup.fitToHeight = 0

    def save_report(self):
        temporary_path = self.output_path + ".tmp"
        try:
            self.wb.save(temporary_path)
            os.replace(temporary_path, self.output_path)
        finally:
            if os.path.exists(temporary_path):
                os.remove(temporary_path)

    def close(self):
        """Cierra el libro de trabajo de openpyxl liberando recursos."""
        if hasattr(self, 'wb') and self.wb is not None:
            try:
                self.wb.close()
            except Exception:
                pass

    def __enter__(self):
        return self

    def __exit__(self, exc_type, exc_val, exc_tb):
        self.close()

    def write_shifts(self, shifts, exceptions, year, month, manual_changes=None):
        meses = ["ENERO", "FEBRERO", "MARZO", "ABRIL", "MAYO", "JUNIO", "JULIO", "AGOSTO", "SEPTIEMBRE", "OCTUBRE", "NOVIEMBRE", "DICIEMBRE"]
        nombre_mes = meses[month - 1]
        
        # Estilos mejorados
        red_fill = PatternFill(start_color="FF3B30", end_color="FF3B30", fill_type="solid")
        da_fill = PatternFill(start_color="B45309", end_color="B45309", fill_type="solid")
        fl_fill = PatternFill(start_color="5B21B6", end_color="5B21B6", fill_type="solid")
        lic_fill = PatternFill(start_color="0E7490", end_color="0E7490", fill_type="solid")
        otr_fill = PatternFill(start_color="374151", end_color="374151", fill_type="solid")
        for_fill = PatternFill(start_color="059669", end_color="059669", fill_type="solid")
        center_align = Alignment(horizontal="center", vertical="center")
        thin_border = Border(
            left=Side(style='thin', color="000000"),
            right=Side(style='thin', color="000000"),
            top=Side(style='thin', color="000000"),
            bottom=Side(style='thin', color="000000")
        )
        bold_font = Font(bold=True)

        # Buscamos la fila de los días (row=2 por defecto en template generado)
        days_row_index = 2
        r_vals = [str(cell.value).strip() if cell.value is not None else "" for cell in self.sheet[days_row_index]]
        if not (any(v in ("1", "1.0", "01") for v in r_vals) and any(v in ("2", "2.0", "02") for v in r_vals)):
            days_row_index = None
            for i, row in enumerate(self.sheet.iter_rows(values_only=True), 1):
                row_vals = [str(val).strip() if val is not None else "" for val in row]
                if any(val in ("1", "1.0", "01") for val in row_vals) and \
                   any(val in ("2", "2.0", "02") for val in row_vals) and \
                   any(val in ("3", "3.0", "03") for val in row_vals):
                    days_row_index = i
                    break
                
        if not days_row_index:
            raise Exception("No se encontró la fila con los días del mes (1, 2, 3...)")
            
        # Determinar en qué columna está cada día del mes
        day_columns = {}
        for col_idx, cell in enumerate(self.sheet[days_row_index], 1):
            val = cell.value
            if val is not None:
                try:
                    val_int = int(float(str(val).strip()))
                    if 1 <= val_int <= 31:
                        if val_int not in day_columns:
                            day_columns[val_int] = col_idx
                except ValueError:
                    pass

        # 0. Combinar y formatear título principal
        self.sheet.merge_cells("A1:AF1")
        title_cell = self.sheet["A1"]
        title_cell.value = f"PLANIFICACIÓN DE TURNOS — {nombre_mes} {year}"
        title_cell.font = Font(bold=True, size=13, color="1F2937")
        title_cell.alignment = Alignment(horizontal="center", vertical="center")
        title_cell.fill = PatternFill(start_color="D9E2F3", end_color="D9E2F3", fill_type="solid")
        self.sheet.row_dimensions[1].height = 28

        # Encontrar en qué fila está cada persona.
        names_col_idx = min(day_columns.values()) - 1
        if names_col_idx < 1:
            names_col_idx = 1
            
        person_rows = {}
        for row_idx in range(days_row_index + 1, self.sheet.max_row + 1):
            name = self.sheet.cell(row=row_idx, column=names_col_idx).value
            if name:
                clean_name = str(name).strip()
                if clean_name.upper() not in ("FUNCIONARIO", "DIA", "DIAS"):
                    person_rows[clean_name] = row_idx

        # 1. LIMPIAR LA GRILLA Y DIBUJAR FINES DE SEMANA DINAMICOS
        gray_fill = openpyxl.styles.PatternFill(start_color="D9D9D9", end_color="D9D9D9", fill_type="solid")
        white_fill = openpyxl.styles.PatternFill(fill_type=None)
        invalid_day_fill = openpyxl.styles.PatternFill(start_color="595959", end_color="595959", fill_type="solid")
        invalid_font = Font(bold=True, color="FFFFFF")
        
        _, days_in_month = calendar.monthrange(year, month)
        dias_letras = ["L", "M", "X", "J", "V", "S", "D"]

        for day_val, col_idx in day_columns.items():
            if day_val > days_in_month:
                # Día no perteneciente al mes (ej: 29-31 en Feb, o 31 en Sep)
                header_cell = self.sheet.cell(row=days_row_index, column=col_idx)
                header_cell.fill = invalid_day_fill
                header_cell.value = "-"
                header_cell.font = invalid_font
                header_cell.alignment = center_align
                weekday_cell = self.sheet.cell(row=days_row_index + 1, column=col_idx)
                weekday_cell.fill = invalid_day_fill
                weekday_cell.value = "-"
                weekday_cell.font = invalid_font
                weekday_cell.alignment = center_align
                for row_idx in person_rows.values():
                    cell = self.sheet.cell(row=row_idx, column=col_idx)
                    cell.value = None
                    cell.fill = invalid_day_fill
                continue

            # Determinar si es fin de semana
            is_weekend = False
            weekday_letter = ""
            try:
                date_obj = datetime.date(year, month, day_val)
                weekday_idx = date_obj.weekday()
                weekday_letter = dias_letras[weekday_idx]
                if weekday_idx >= 5: # 5=Sat, 6=Sun
                    is_weekend = True
            except ValueError:
                pass
                
            fill_to_apply = gray_fill if is_weekend else white_fill
            
            # Pintar el encabezado de día y de letra de semana (fila 2 y 3)
            header_cell = self.sheet.cell(row=days_row_index, column=col_idx)
            header_cell.fill = fill_to_apply
            header_cell.value = day_val

            weekday_cell = self.sheet.cell(row=days_row_index + 1, column=col_idx)
            weekday_cell.fill = fill_to_apply
            weekday_cell.value = weekday_letter
            weekday_cell.alignment = center_align
            weekday_cell.font = bold_font
            
            # Pintar las filas del personal y vaciar texto
            for row_idx in person_rows.values():
                cell = self.sheet.cell(row=row_idx, column=col_idx)
                cell.value = None
                cell.fill = fill_to_apply

        # Preparar mapa de cambios manuales
        if manual_changes is None:
            manual_changes = []
            for sh in shifts:
                if sh.get('es_manual') or sh.get('es_forzado'):
                    manual_changes.append({
                        'semana': sh.get('semana'),
                        'nuevo': sh.get('persona', 'Sin asignar'),
                        'anterior': sh.get('anterior', 'Rotación automática'),
                        'motivo': sh.get('motivo', 'Cambio manual registrado')
                    })

        manual_map = {mc['semana']: mc for mc in manual_changes if mc.get('semana')}

        # Escribir los turnos
        for shift in shifts:
            person = shift.get('persona')
            if not person or person not in person_rows:
                continue
            
            row_idx = person_rows[person]
            start_date, end_date = shift['semana']
            
            mc_info = manual_map.get(shift['semana'])
            is_manual = bool(shift.get('es_manual') or shift.get('es_forzado') or mc_info)
            shift_fill = for_fill if is_manual else red_fill
            
            comment_text = None
            if is_manual:
                orig_p = (mc_info.get('anterior') if mc_info else None) or shift.get('anterior') or "Rotación automática"
                mot = (mc_info.get('motivo') if mc_info else None) or shift.get('motivo') or "Cambio manual registrado"
                comment_text = (
                    f"CAMBIO MANUAL DE GUARDIA\n"
                    f"Asignado: {person}\n"
                    f"Guardia original: {orig_p}\n"
                    f"Motivo: {mot}"
                )

            current_day = start_date
            while current_day <= end_date:
                if current_day.month == month and current_day.year == year:
                    day_val = current_day.day
                    if day_val in day_columns:
                        col_idx = day_columns[day_val]
                        cell = self.sheet.cell(row=row_idx, column=col_idx)
                        cell.fill = shift_fill
                        cell.border = thin_border
                        if comment_text:
                            cell.comment = Comment(comment_text, "Sistema de Turnos")
                current_day += timedelta(days=1)

        # Escribir las excepciones (DA, FL)
        for exc in exceptions:
            person = exc['persona']
            exc_date = exc['fecha']
            
            if person in person_rows and exc_date.month == month and exc_date.year == year:
                row_idx = person_rows[person]
                day_val = exc_date.day
                if day_val in day_columns:
                    col_idx = day_columns[day_val]
                    cell = self.sheet.cell(row=row_idx, column=col_idx)
                    cell.value = exc['tipo']
                    _exc_fills = {'DA': da_fill, 'FL': fl_fill, 'LIC': lic_fill, 'OTR': otr_fill, 'FOR': for_fill}
                    cell.fill = _exc_fills.get(exc['tipo'], otr_fill)
                    cell.alignment = center_align
                    cell.font = bold_font
                    cell.border = thin_border
                    if exc.get('tipo') == 'OTR' and exc.get('motivo'):
                        cell.comment = Comment(f"Motivo: {exc['motivo']}", "Sistema de Turnos")

        # 3. ESCRIBIR TOTALES POR FUNCIONARIO (COLUMNAS 33 a 37: AG, AH, AI, AJ, AK)
        summary_fill = PatternFill(start_color="F2F5F9", end_color="F2F5F9", fill_type="solid")
        summary_font = Font(bold=True, size=10)

        for person, row_idx in person_rows.items():
            turnos_cnt = 0
            for sh in shifts:
                if sh.get('persona') == person:
                    s_d, e_d = sh['semana']
                    if (s_d.year == year and s_d.month == month) or (e_d.year == year and e_d.month == month):
                        turnos_cnt += 1

            da_cnt = 0
            fl_cnt = 0
            lic_cnt = 0
            otr_cnt = 0
            for exc in exceptions:
                if exc.get('persona') == person:
                    f = exc.get('fecha')
                    if f and f.year == year and f.month == month:
                        tipo = exc.get('tipo')
                        if tipo == 'DA':
                            da_cnt += 1
                        elif tipo == 'FL':
                            fl_cnt += 1
                        elif tipo == 'LIC':
                            lic_cnt += 1
                        elif tipo in ('OTR', 'FOR'):
                            otr_cnt += 1

            totals = [
                (33, turnos_cnt),
                (34, da_cnt),
                (35, fl_cnt),
                (36, lic_cnt),
                (37, otr_cnt),
            ]
            for col_idx, val in totals:
                tot_cell = self.sheet.cell(row=row_idx, column=col_idx, value=val)
                tot_cell.alignment = center_align
                tot_cell.font = summary_font
                tot_cell.fill = summary_fill
                tot_cell.border = thin_border

        # 4. DIBUJAR LEYENDA EXPLICATIVA DE COLORES
        max_person_row = max(person_rows.values()) if person_rows else (days_row_index + len(self.personal))
        legend_title_row = max_person_row + 3

        title_lbl = self.sheet.cell(row=legend_title_row, column=1, value="CONVENCIONES Y LEYENDA")
        title_lbl.font = Font(bold=True, size=10, color="1F2937")
        self.sheet.row_dimensions[legend_title_row].height = 22

        left_items = [
            ("Turno de Guardia", red_fill, "■", "FFFFFF"),
            ("DA: Día Administrativo", da_fill, "DA", "FFFFFF"),
            ("FL: Feriado Legal", fl_fill, "FL", "FFFFFF"),
            ("LIC: Licencia Médica", lic_fill, "LIC", "FFFFFF"),
        ]
        right_items = [
            ("OTR: Otro Permiso", otr_fill, "OTR", "FFFFFF"),
            ("Fin de Semana", gray_fill, " ", "000000"),
            ("Cambio Guardia (Manual)", for_fill, "■", "FFFFFF"),
        ]

        for idx, (label, fill_style, mark, text_color) in enumerate(left_items):
            r = legend_title_row + 1 + idx
            self.sheet.row_dimensions[r].height = 20
            # Muestra de color
            chip = self.sheet.cell(row=r, column=2, value=mark)
            chip.fill = fill_style
            chip.alignment = center_align
            chip.font = Font(bold=True, size=9, color=text_color)
            chip.border = thin_border
            # Texto explicativo
            self.sheet.merge_cells(start_row=r, start_column=3, end_row=r, end_column=8)
            desc = self.sheet.cell(row=r, column=3, value=label)
            desc.font = Font(size=9.5, color="374151")
            desc.alignment = Alignment(horizontal="left", vertical="center")

        for idx, (label, fill_style, mark, text_color) in enumerate(right_items):
            r = legend_title_row + 1 + idx
            # Muestra de color
            chip = self.sheet.cell(row=r, column=10, value=mark)
            chip.fill = fill_style
            chip.alignment = center_align
            chip.font = Font(bold=True, size=9, color=text_color)
            chip.border = thin_border
            # Texto explicativo
            self.sheet.merge_cells(start_row=r, start_column=11, end_row=r, end_column=17)
            desc = self.sheet.cell(row=r, column=11, value=label)
            desc.font = Font(size=9.5, color="374151")
            desc.alignment = Alignment(horizontal="left", vertical="center")

        # 5. DIBUJAR TABLA DE REGISTRO DE CAMBIOS MANUALES (SI EXISTEN EN EL MES)
        next_table_row = legend_title_row + max(len(left_items), len(right_items)) + 2

        if manual_changes:
            mc_title = self.sheet.cell(
                row=next_table_row, column=1,
                value="REGISTRO DE CAMBIOS MANUALES DE GUARDIA (PERMUTAS / ACCIDENTES)"
            )
            mc_title.font = Font(bold=True, size=10, color="065F46")
            self.sheet.row_dimensions[next_table_row].height = 22

            hdr_row = next_table_row + 1
            self.sheet.row_dimensions[hdr_row].height = 20

            hdr_fill_mc = PatternFill(start_color="D1FAE5", end_color="D1FAE5", fill_type="solid")
            hdr_font_mc = Font(bold=True, size=9, color="065F46")

            # Col 1: FUNCIONARIO ASIGNADO
            c1 = self.sheet.cell(row=hdr_row, column=1, value="FUNCIONARIO ASIGNADO")
            c1.fill = hdr_fill_mc
            c1.font = hdr_font_mc
            c1.alignment = Alignment(horizontal="left", vertical="center")
            c1.border = thin_border

            # Cols 2-5: SEMANA
            self.sheet.merge_cells(start_row=hdr_row, start_column=2, end_row=hdr_row, end_column=5)
            for c in range(2, 6):
                ch = self.sheet.cell(row=hdr_row, column=c)
                ch.fill = hdr_fill_mc
                ch.font = hdr_font_mc
                ch.border = thin_border
            c2 = self.sheet.cell(row=hdr_row, column=2, value="SEMANA")
            c2.alignment = center_align

            # Cols 6-11: GUARDIA ORIGINAL
            self.sheet.merge_cells(start_row=hdr_row, start_column=6, end_row=hdr_row, end_column=11)
            for c in range(6, 12):
                ch = self.sheet.cell(row=hdr_row, column=c)
                ch.fill = hdr_fill_mc
                ch.font = hdr_font_mc
                ch.border = thin_border
            c3 = self.sheet.cell(row=hdr_row, column=6, value="GUARDIA ORIGINAL")
            c3.alignment = Alignment(horizontal="left", vertical="center")

            # Cols 12-26: MOTIVO DEL CAMBIO / JUSTIFICACIÓN
            self.sheet.merge_cells(start_row=hdr_row, start_column=12, end_row=hdr_row, end_column=26)
            for c in range(12, 27):
                ch = self.sheet.cell(row=hdr_row, column=c)
                ch.fill = hdr_fill_mc
                ch.font = hdr_font_mc
                ch.border = thin_border
            c4 = self.sheet.cell(row=hdr_row, column=12, value="MOTIVO DEL CAMBIO / JUSTIFICACIÓN")
            c4.alignment = Alignment(horizontal="left", vertical="center")

            # Filas de datos
            data_font = Font(size=9.5, color="374151")
            for idx, mc in enumerate(manual_changes):
                curr_r = hdr_row + 1 + idx
                self.sheet.row_dimensions[curr_r].height = 20

                # Funcionario Asignado
                p_cell = self.sheet.cell(row=curr_r, column=1, value=mc.get('nuevo', 'Sin asignar'))
                p_cell.font = data_font
                p_cell.alignment = Alignment(horizontal="left", vertical="center")
                p_cell.border = thin_border

                # Semana
                s_d, e_d = mc.get('semana', (None, None))
                if s_d and e_d:
                    semana_str = f"{s_d.strftime('%d/%m')} al {e_d.strftime('%d/%m')}"
                else:
                    semana_str = "Semana N/A"
                self.sheet.merge_cells(start_row=curr_r, start_column=2, end_row=curr_r, end_column=5)
                for c in range(2, 6):
                    self.sheet.cell(row=curr_r, column=c).border = thin_border
                d_cell = self.sheet.cell(row=curr_r, column=2, value=semana_str)
                d_cell.font = data_font
                d_cell.alignment = center_align

                # Guardia Original
                self.sheet.merge_cells(start_row=curr_r, start_column=6, end_row=curr_r, end_column=11)
                for c in range(6, 12):
                    self.sheet.cell(row=curr_r, column=c).border = thin_border
                orig_cell = self.sheet.cell(row=curr_r, column=6, value=mc.get('anterior', 'Rotación automática'))
                orig_cell.font = data_font
                orig_cell.alignment = Alignment(horizontal="left", vertical="center")

                # Motivo
                self.sheet.merge_cells(start_row=curr_r, start_column=12, end_row=curr_r, end_column=26)
                for c in range(12, 27):
                    self.sheet.cell(row=curr_r, column=c).border = thin_border
                m_cell = self.sheet.cell(row=curr_r, column=12, value=mc.get('motivo', 'Sin justificación'))
                m_cell.font = data_font
                m_cell.alignment = Alignment(horizontal="left", vertical="center")

            next_table_row = hdr_row + len(manual_changes) + 2

        # 6. DIBUJAR TABLA DE DETALLE DE EXCEPCIONES OTR (SI EXISTEN EN EL MES)
        otr_excs = [
            e for e in exceptions
            if e.get('tipo') == 'OTR' and e.get('fecha') and
            e['fecha'].year == year and e['fecha'].month == month
        ]

        if otr_excs:
            otr_start_row = next_table_row

            otr_title = self.sheet.cell(row=otr_start_row, column=1, value="DETALLE DE PERMISOS Y EXCEPCIONES ESPECIALES (OTR)")
            otr_title.font = Font(bold=True, size=10, color="1F2937")
            self.sheet.row_dimensions[otr_start_row].height = 22

            # Agrupar por (persona, motivo) para presentar rangos continuos
            groups = {}
            for e in sorted(otr_excs, key=lambda x: (x.get('persona', ''), x.get('fecha'))):
                p = e.get('persona', '')
                m = (e.get('motivo') or 'Sin justificación especificada').strip()
                k = (p, m)
                if k not in groups:
                    groups[k] = []
                groups[k].append(e['fecha'])

            # Encabezados de la tabla OTR
            hdr_row = otr_start_row + 1
            self.sheet.row_dimensions[hdr_row].height = 20

            hdr_fill = PatternFill(start_color="D9E2F3", end_color="D9E2F3", fill_type="solid")
            hdr_font = Font(bold=True, size=9)

            c1 = self.sheet.cell(row=hdr_row, column=1, value="FUNCIONARIO")
            c1.fill = hdr_fill
            c1.font = hdr_font
            c1.alignment = Alignment(horizontal="left", vertical="center")
            c1.border = thin_border

            self.sheet.merge_cells(start_row=hdr_row, start_column=2, end_row=hdr_row, end_column=5)
            for c in range(2, 6):
                cell_h = self.sheet.cell(row=hdr_row, column=c)
                cell_h.fill = hdr_fill
                cell_h.font = hdr_font
                cell_h.border = thin_border
            c2 = self.sheet.cell(row=hdr_row, column=2, value="DÍAS / PERÍODO")
            c2.alignment = center_align

            self.sheet.merge_cells(start_row=hdr_row, start_column=6, end_row=hdr_row, end_column=22)
            for c in range(6, 23):
                cell_h = self.sheet.cell(row=hdr_row, column=c)
                cell_h.fill = hdr_fill
                cell_h.font = hdr_font
                cell_h.border = thin_border
            c3 = self.sheet.cell(row=hdr_row, column=6, value="MOTIVO / JUSTIFICACIÓN")
            c3.alignment = Alignment(horizontal="left", vertical="center")

            # Filas de datos
            data_font = Font(size=9.5, color="374151")
            for idx, ((person, motivo), dates) in enumerate(groups.items()):
                curr_r = hdr_row + 1 + idx
                self.sheet.row_dimensions[curr_r].height = 20

                # Funcionario
                p_cell = self.sheet.cell(row=curr_r, column=1, value=person)
                p_cell.font = data_font
                p_cell.alignment = Alignment(horizontal="left", vertical="center")
                p_cell.border = thin_border

                # Días formateados
                dates_str = self._format_date_ranges(dates)
                self.sheet.merge_cells(start_row=curr_r, start_column=2, end_row=curr_r, end_column=5)
                for c in range(2, 6):
                    self.sheet.cell(row=curr_r, column=c).border = thin_border
                d_cell = self.sheet.cell(row=curr_r, column=2, value=dates_str)
                d_cell.font = data_font
                d_cell.alignment = center_align

                # Motivo
                self.sheet.merge_cells(start_row=curr_r, start_column=6, end_row=curr_r, end_column=22)
                for c in range(6, 23):
                    self.sheet.cell(row=curr_r, column=c).border = thin_border
                m_cell = self.sheet.cell(row=curr_r, column=6, value=motivo)
                m_cell.font = data_font
                m_cell.alignment = Alignment(horizontal="left", vertical="center")

    @staticmethod
    def _format_date_ranges(dates):
        """Agrupa fechas consecutivas en formato legible (ej: '01/09 al 04/09, 10/09')."""
        if not dates:
            return ""
        sorted_dates = sorted(dates)
        ranges = []
        start = sorted_dates[0]
        prev = sorted_dates[0]
        for d in sorted_dates[1:]:
            if d == prev + timedelta(days=1):
                prev = d
            else:
                if start == prev:
                    ranges.append(start.strftime("%d/%m"))
                else:
                    ranges.append(f"{start.strftime('%d/%m')} al {prev.strftime('%d/%m')}")
                start = d
                prev = d
        if start == prev:
            ranges.append(start.strftime("%d/%m"))
        else:
            ranges.append(f"{start.strftime('%d/%m')} al {prev.strftime('%d/%m')}")
        return ", ".join(ranges)
