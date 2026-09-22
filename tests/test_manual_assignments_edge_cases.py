import unittest
import os
import json
import tempfile
from datetime import date
from models.shift_manager import ShiftManager
from controllers.main_controller import MainController
from utils.excel_handler import ExcelHandler

class TestManualAssignmentsEdgeCases(unittest.TestCase):
    def setUp(self):
        self.temp_dir = tempfile.TemporaryDirectory()
        self.config_path = os.path.join(self.temp_dir.name, "config.json")
        self.output_excel = os.path.join(self.temp_dir.name, "turnos_generados.xlsx")
        self.initial_personal = [
            {"id": 1, "nombre": "SGT (F) PEREZ JUAN"},
            {"id": 2, "nombre": "CBO (M) GOMEZ ANA"},
            {"id": 3, "nombre": "CBO (F) DIAZ LUIS"},
            {"id": 4, "nombre": "SBC (M) SILVA MARIA"},
            {"id": 5, "nombre": "CBO (F) ROJAS CARLOS"},
            {"id": 6, "nombre": "SGT (M) SOTO ANDRES"},
            {"id": 7, "nombre": "CBO (M) VERA ELENA"},
            {"id": 8, "nombre": "SBC (F) MORA PATRICIA"}
        ]
        self.initial_data = {
            "personal": self.initial_personal,
            "inicio": {},
            "historial": {},
            "siguiente_id": 1,
            "pendientes": [],
            "snapshots": {},
            "excepciones": {},
            "asignaciones_manuales": {}
        }
        with open(self.config_path, "w", encoding="utf-8") as f:
            json.dump(self.initial_data, f)
        self.controller = MainController(self.temp_dir.name)
        self.manager = self.controller.shift_manager

    def tearDown(self):
        self.temp_dir.cleanup()

    def test_past_month_saved_with_manual_assignment_preserves_current_queue(self):
        # 1. Avanzar frente operativo: Guardar Julio y Agosto 2026
        self.controller.advance_queue(2026, 7, [])
        self.controller.advance_queue(2026, 8, [])

        present_id = self.manager.siguiente_id
        present_pendientes = list(self.manager.pendientes)

        # 2. Guardar retroactivamente Junio 2026 (mes anterior al frente) con asignación manual
        jun_shifts = self.controller.preview_shifts(2026, 6, [])
        first_week_key = f"{jun_shifts[0]['semana'][0].strftime('%Y-%m-%d')}_{jun_shifts[0]['semana'][1].strftime('%Y-%m-%d')}"
        manual_assignment = {first_week_key: "SBC (M) SILVA MARIA"}

        ok, _ = self.controller.advance_queue(2026, 6, [], manual_assignments=manual_assignment)
        self.assertTrue(ok)

        # El puntero y la cola operativa del presente NO se deben haber alterado
        self.assertEqual(self.manager.siguiente_id, present_id)
        self.assertEqual(self.manager.pendientes, present_pendientes)

        # El historial de Junio debe tener a Silva María en la semana 1
        self.assertEqual(self.manager.historial[first_week_key], "SBC (M) SILVA MARIA")
        # Las asignaciones manuales de Junio deben haberse persistido
        saved_manuals = self.manager.get_manual_assignments("2026-06")
        self.assertEqual(saved_manuals.get(first_week_key), "SBC (M) SILVA MARIA")

    def test_past_month_shift_retains_es_manual_flag_when_viewed_later(self):
        # Guardar mes con asignación manual
        shifts = self.controller.preview_shifts(2026, 5, [])
        wk = f"{shifts[1]['semana'][0].strftime('%Y-%m-%d')}_{shifts[1]['semana'][1].strftime('%Y-%m-%d')}"
        self.controller.advance_queue(2026, 5, [], manual_assignments={wk: "CBO (M) VERA ELENA"})

        # Avanzar a junio para dejar a mayo en el pasado
        self.controller.advance_queue(2026, 6, [])

        # Consultar mayo nuevamente (como lo hace el Calendario o TabPlan al ver historial)
        loaded_shifts = self.controller.preview_shifts(2026, 5, [])
        target_shift = next(sh for sh in loaded_shifts if sh['semana'] == shifts[1]['semana'])

        self.assertEqual(target_shift['persona'], "CBO (M) VERA ELENA")
        self.assertTrue(target_shift.get('es_manual'), "El turno histórico debe conservar la marca es_manual")

    def test_editing_manual_assignment_in_already_closed_month(self):
        # 1. Guardar mes inicialmente con Gomez Ana en la primera semana
        shifts = self.controller.preview_shifts(2026, 4, [])
        wk = f"{shifts[0]['semana'][0].strftime('%Y-%m-%d')}_{shifts[0]['semana'][1].strftime('%Y-%m-%d')}"
        self.controller.advance_queue(2026, 4, [], manual_assignments={wk: "CBO (M) GOMEZ ANA"})
        self.assertEqual(self.manager.historial[wk], "CBO (M) GOMEZ ANA")

        # 2. Re-abrir Abril y cambiar la asignación manual a Soto Andres
        re_shifts = self.controller.preview_shifts(2026, 4, [], manual_assignments={wk: "SGT (M) SOTO ANDRES"})
        self.assertEqual(re_shifts[0]['persona'], "SGT (M) SOTO ANDRES")
        self.assertTrue(re_shifts[0].get('es_manual'))

        # 3. Guardar mes editado
        ok, _ = self.controller.advance_queue(2026, 4, [], manual_assignments={wk: "SGT (M) SOTO ANDRES"})
        self.assertTrue(ok)
        self.assertEqual(self.manager.historial[wk], "SGT (M) SOTO ANDRES")
        self.assertEqual(self.manager.get_manual_assignments("2026-04")[wk], "SGT (M) SOTO ANDRES")

    def test_manual_assignment_consecutive_weeks_same_person(self):
        # Asignar a la misma persona en semanas consecutivas 0 y 1
        shifts = self.controller.preview_shifts(2026, 10, [])
        wk0 = f"{shifts[0]['semana'][0].strftime('%Y-%m-%d')}_{shifts[0]['semana'][1].strftime('%Y-%m-%d')}"
        wk1 = f"{shifts[1]['semana'][0].strftime('%Y-%m-%d')}_{shifts[1]['semana'][1].strftime('%Y-%m-%d')}"

        manual_dict = {
            wk0: "SBC (F) MORA PATRICIA",
            wk1: "SBC (F) MORA PATRICIA"
        }
        res_shifts = self.controller.preview_shifts(2026, 10, [], manual_assignments=manual_dict)

        self.assertEqual(res_shifts[0]['persona'], "SBC (F) MORA PATRICIA")
        self.assertTrue(res_shifts[0].get('es_manual'))
        self.assertEqual(res_shifts[1]['persona'], "SBC (F) MORA PATRICIA")
        self.assertTrue(res_shifts[1].get('es_manual'))

        # Semanas siguientes continúan con el resto del personal sin duplicados vacíos
        for sh in res_shifts[2:]:
            self.assertIsNotNone(sh['persona'])

    def test_manual_assignment_when_person_has_exception_in_same_week(self):
        # Perez Juan tiene DA el día 6 de Octubre (semana 1)
        shifts = self.controller.preview_shifts(2026, 10, [])
        wk0 = f"{shifts[0]['semana'][0].strftime('%Y-%m-%d')}_{shifts[0]['semana'][1].strftime('%Y-%m-%d')}"
        da_exc = [{'persona': 'SGT (F) PEREZ JUAN', 'fecha': date(2026, 10, 6), 'tipo': 'DA'}]

        # Se fuerza manualmente a Perez Juan en esa misma semana
        res_shifts = self.controller.preview_shifts(
            2026, 10, da_exc, manual_assignments={wk0: "SGT (F) PEREZ JUAN"}
        )

        # La semana queda asignada a Perez Juan con flag es_manual
        self.assertEqual(res_shifts[0]['persona'], "SGT (F) PEREZ JUAN")
        self.assertTrue(res_shifts[0].get('es_manual'))

    def test_manual_assignment_of_person_in_pendientes_clears_from_pendientes(self):
        # Colocar a Diaz Luis (id 3) en pendientes
        self.manager.pendientes = [3]
        self.manager.siguiente_id = 1
        self.manager.save_config()

        # En la primera semana asignamos manualmente a Diaz Luis
        shifts = self.controller.preview_shifts(2026, 10, [])
        wk0 = f"{shifts[0]['semana'][0].strftime('%Y-%m-%d')}_{shifts[0]['semana'][1].strftime('%Y-%m-%d')}"

        # Al avanzar mes con Diaz Luis asignado manualmente
        self.controller.advance_queue(2026, 10, [], manual_assignments={wk0: "CBO (F) DIAZ LUIS"})

        # Diaz Luis ya no debe estar en pendientes
        self.assertNotIn(3, self.manager.pendientes)

    def test_manual_assignment_of_other_person_keeps_pendientes_intact(self):
        # Configurar snapshot al inicio de 2026-10 con Gomez Ana (id 2) en pendientes
        self.manager.snapshots["2026-10"] = {"siguiente_id": 1, "pendientes": [2]}
        self.manager.siguiente_id = 1
        self.manager.save_config()

        # Forzamos manualmente a Mora Patricia (id 8) en la semana 1
        shifts = self.controller.preview_shifts(2026, 10, [])
        wk0 = f"{shifts[0]['semana'][0].strftime('%Y-%m-%d')}_{shifts[0]['semana'][1].strftime('%Y-%m-%d')}"
        manual_dict = {wk0: "SBC (F) MORA PATRICIA"}

        gen_shifts = self.controller.preview_shifts(2026, 10, [], manual_assignments=manual_dict)
        # Semana 1 es Mora Patricia (manual)
        self.assertEqual(gen_shifts[0]['persona'], "SBC (F) MORA PATRICIA")
        # Semana 2 debe recuperar Gomez Ana que estaba en pendientes
        self.assertEqual(gen_shifts[1]['persona'], "CBO (M) GOMEZ ANA")

    def test_excel_export_parity_with_manual_assignments(self):
        # Generar turnos para 2026-10 con asignación manual
        shifts = self.controller.preview_shifts(2026, 10, [])
        wk0 = f"{shifts[0]['semana'][0].strftime('%Y-%m-%d')}_{shifts[0]['semana'][1].strftime('%Y-%m-%d')}"
        manual_dict = {wk0: "SBC (M) SILVA MARIA"}

        target_file = os.path.join(self.temp_dir.name, "excel_test.xlsx")
        ok, msg = self.controller.process_generation(
            2026, 10, [], manual_assignments=manual_dict, target_path=target_file
        )
        self.assertTrue(ok, msg)
        self.assertTrue(os.path.exists(target_file))

        # Verificar que el Excel se abre y contiene los turnos de Silva Maria
        import openpyxl
        wb = openpyxl.load_workbook(target_file)
        sheet = wb.active

        # Buscar fila de SILVA MARIA
        silva_row = None
        for r in range(1, sheet.max_row + 1):
            if "SILVA MARIA" in str(sheet.cell(row=r, column=1).value):
                silva_row = r
                break
        self.assertIsNotNone(silva_row)

        # En la primera semana (días correspondientes a octubre), debe haber celdas pintadas de rojo
        s_date, e_date = shifts[0]['semana']
        cur = s_date
        has_red_cell = False
        while cur <= e_date:
            if cur.month == 10 and cur.year == 2026:
                day_col = cur.day + 1
                cell_fill = sheet.cell(row=silva_row, column=day_col).fill
                if cell_fill and cell_fill.start_color and cell_fill.start_color.rgb:
                    # Color rojo del turno
                    if "FF3B30" in str(cell_fill.start_color.rgb).upper() or "3B30" in str(cell_fill.start_color.rgb).upper():
                        has_red_cell = True
            cur += date.resolution
        self.assertTrue(has_red_cell, "El funcionario asignado manualmente debe tener celdas de turno en Excel")

    def test_manual_assignment_persists_even_if_person_is_deleted_later(self):
        # Asignar a Vera Elena en 2026-09 y cerrar mes
        shifts = self.controller.preview_shifts(2026, 9, [])
        wk0 = f"{shifts[0]['semana'][0].strftime('%Y-%m-%d')}_{shifts[0]['semana'][1].strftime('%Y-%m-%d')}"
        self.controller.advance_queue(2026, 9, [], manual_assignments={wk0: "CBO (M) VERA ELENA"})

        # Eliminar a Vera Elena (id 7) del personal
        ok_del, _ = self.controller.remove_person(7)
        self.assertTrue(ok_del)

        # El historial y las asignaciones manuales deben conservar el registro
        saved_manuals = self.manager.get_manual_assignments("2026-09")
        self.assertEqual(saved_manuals[wk0], "CBO (M) VERA ELENA")
        self.assertEqual(self.manager.historial[wk0], "CBO (M) VERA ELENA")

        # Al previsualizar el mes 2026-09, Vera Elena sigue en el turno de esa semana
        past_shifts = self.controller.preview_shifts(2026, 9, [])
        self.assertEqual(past_shifts[0]['persona'], "CBO (M) VERA ELENA")
        self.assertTrue(past_shifts[0].get('es_manual'))

if __name__ == '__main__':
    unittest.main()
