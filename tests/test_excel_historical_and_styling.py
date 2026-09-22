import json
import os
import tempfile
import unittest
from datetime import date
from pathlib import Path

import openpyxl

from controllers.main_controller import MainController
from utils.excel_handler import ExcelHandler


class ExcelHistoricalAndStylingTests(unittest.TestCase):
    def setUp(self):
        self.temp_dir = tempfile.TemporaryDirectory()
        self.root_path = Path(self.temp_dir.name)
        self.config_path = self.root_path / "config.json"

        # Configuración inicial con 3 personas
        self.personal = [
            {"id": 1, "nombre": "COM ALARCON JOSE"},
            {"id": 2, "nombre": "PRO BRAVO PATRICIO"},
            {"id": 3, "nombre": "SBC CARVAJAL LUIS"},
        ]
        self.config_data = {
            "personal": self.personal,
            "inicio": {},
            "historial": {
                # Semana en septiembre 2026 asignada a COM ALARCON JOSE
                "2026-09-07_2026-09-13": "COM ALARCON JOSE",
                # Semana en septiembre 2026 asignada a SBC CARVAJAL LUIS
                "2026-09-14_2026-09-20": "SBC CARVAJAL LUIS",
            },
            "siguiente_id": 2,
            "pendientes": [],
            "snapshots": {},
            "excepciones": {
                "2026-09": [
                    {"persona": "PRO BRAVO PATRICIO", "fecha": "2026-09-15", "tipo": "DA"}
                ]
            },
        }
        self.config_path.write_text(
            json.dumps(self.config_data, indent=2, ensure_ascii=False), encoding="utf-8"
        )
        self.controller = MainController(str(self.root_path))

    def tearDown(self):
        self.temp_dir.cleanup()

    def test_deleted_personnel_included_in_excel_export(self):
        # 1. Eliminar a SBC CARVAJAL LUIS de la lista activa de personal
        success, _ = self.controller.remove_person(person_id=3)
        self.assertTrue(success)

        # Constatar que ya no está en el personal activo
        active_names = [p["nombre"] for p in self.controller.get_all_persons()]
        self.assertNotIn("SBC CARVAJAL LUIS", active_names)

        # 2. Exportar el mes de Septiembre 2026 donde Carvajal tiene un turno en historial
        output_file = self.root_path / "test_deleted_person_export.xlsx"
        ok, msg = self.controller.process_generation(
            2026, 9, [], target_path=str(output_file)
        )
        self.assertTrue(ok, f"Fallo al exportar: {msg}")
        self.assertTrue(output_file.exists())

        # 3. Abrir el archivo Excel y verificar que SBC CARVAJAL LUIS tiene su fila y su turno marcado
        wb = openpyxl.load_workbook(str(output_file))
        sheet = wb["Turnos"]

        # Buscar la fila de SBC CARVAJAL LUIS
        carvajal_row = None
        for r in range(4, sheet.max_row + 1):
            if sheet.cell(row=r, column=1).value == "SBC CARVAJAL LUIS":
                carvajal_row = r
                break

        self.assertIsNotNone(
            carvajal_row,
            "SBC CARVAJAL LUIS debió ser incluido en el reporte Excel a pesar de haber sido eliminado de la lista activa.",
        )

        # Verificar que el día 16 de septiembre (miércoles) tiene color de turno (FF3B30)
        # Columna 1=funcionario, Col 2=día 1 ... Col 17=día 16
        day_16_cell = sheet.cell(row=carvajal_row, column=17)
        self.assertIsNotNone(day_16_cell.fill.start_color.rgb)
        self.assertIn("FF3B30", day_16_cell.fill.start_color.rgb)

    def test_excel_title_is_merged_and_formatted(self):
        output_file = self.root_path / "test_title.xlsx"
        ok, msg = self.controller.process_generation(
            2026, 9, [], target_path=str(output_file)
        )
        self.assertTrue(ok, f"Fallo al exportar: {msg}")

        wb = openpyxl.load_workbook(str(output_file))
        sheet = wb["Turnos"]

        # Verificar que el rango A1:AF1 está combinado
        merged_ranges = [str(r) for r in sheet.merged_cells.ranges]
        self.assertIn("A1:AF1", merged_ranges)

        title_cell = sheet["A1"]
        self.assertEqual(title_cell.value, "PLANIFICACIÓN DE TURNOS — SEPTIEMBRE 2026")
        self.assertTrue(title_cell.font.bold)
        self.assertEqual(title_cell.alignment.horizontal, "center")

    def test_excel_legend_block_generated_correctly(self):
        output_file = self.root_path / "test_legend.xlsx"
        ok, msg = self.controller.process_generation(
            2026, 9, [], target_path=str(output_file)
        )
        self.assertTrue(ok, f"Fallo al exportar: {msg}")

        wb = openpyxl.load_workbook(str(output_file))
        sheet = wb["Turnos"]

        # Buscar la celda con el texto de la leyenda
        legend_row = None
        for r in range(5, sheet.max_row + 1):
            if sheet.cell(row=r, column=1).value == "CONVENCIONES Y LEYENDA":
                legend_row = r
                break

        self.assertIsNotNone(legend_row, "La sección 'CONVENCIONES Y LEYENDA' debe existir.")

        # Verificar chips y descripciones en las filas de leyenda
        # Fila 1 de leyenda (Turno de guardia en columna 2, texto en columna 3)
        r1 = legend_row + 1
        chip_turno = sheet.cell(row=r1, column=2)
        desc_turno = sheet.cell(row=r1, column=3)
        self.assertEqual(chip_turno.value, "■")
        self.assertIn("FF3B30", chip_turno.fill.start_color.rgb)
        self.assertEqual(desc_turno.value, "Turno de Guardia")

        # Fila 2 de leyenda (DA en columna 2, texto en columna 3)
        r2 = legend_row + 2
        chip_da = sheet.cell(row=r2, column=2)
        desc_da = sheet.cell(row=r2, column=3)
        self.assertEqual(chip_da.value, "DA")
        self.assertIn("B45309", chip_da.fill.start_color.rgb)
        self.assertEqual(desc_da.value, "DA: Día Administrativo")

        # Verificar columna derecha de la leyenda (columna 10 chip, columna 11 texto)
        chip_for = sheet.cell(row=r2, column=10)
        desc_for = sheet.cell(row=r2, column=11)
        self.assertEqual(chip_for.value, "FOR")
        self.assertIn("059669", chip_for.fill.start_color.rgb)
        self.assertEqual(desc_for.value, "FOR: Asignación Forzada")


if __name__ == "__main__":
    unittest.main()
