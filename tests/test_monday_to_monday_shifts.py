import json
import os
import tempfile
import unittest
from datetime import date
from pathlib import Path
import openpyxl

from controllers.main_controller import MainController
from models.shift_manager import ShiftManager
from utils.excel_handler import ExcelHandler


class TestMondayToMondayShifts(unittest.TestCase):
    def setUp(self):
        self.temp_dir = tempfile.TemporaryDirectory()
        self.root_path = Path(self.temp_dir.name)
        self.config_path = self.root_path / "config.json"

        self.personal = [
            {"id": 1, "nombre": "SGT (F) PEREZ JUAN", "email": "perez@test.com"},
            {"id": 2, "nombre": "CBO (M) GOMEZ ANA", "email": "gomez@test.com"},
            {"id": 3, "nombre": "CBO (F) DIAZ LUIS", "email": "diaz@test.com"},
            {"id": 4, "nombre": "CBO (F) SILVA MARIA", "email": "silva@test.com"},
            {"id": 5, "nombre": "CBO (M) ROJAS PEDRO", "email": "rojas@test.com"},
        ]

        self.initial_data = {
            "personal": self.personal,
            "inicio": {},
            "historial": {
                # Septiembre 2026 cerrado en formato histórico (lunes a domingo)
                "2026-08-31_2026-09-06": "SGT (F) PEREZ JUAN",
                "2026-09-07_2026-09-13": "CBO (M) GOMEZ ANA",
                "2026-09-14_2026-09-20": "CBO (F) DIAZ LUIS",
                "2026-09-21_2026-09-27": "CBO (F) SILVA MARIA",
                "2026-09-28_2026-10-04": "CBO (M) ROJAS PEDRO",
            },
            "siguiente_id": 1,
            "pendientes": [],
            "snapshots": {},
            "excepciones": {},
            "asignaciones_manuales": {},
        }
        self.config_path.write_text(json.dumps(self.initial_data, indent=2, ensure_ascii=False), encoding="utf-8")
        self.controller = MainController(str(self.root_path))
        self.manager = self.controller.shift_manager

    def tearDown(self):
        self.temp_dir.cleanup()

    def test_build_weeks_transition_from_september_to_october(self):
        """Verifica que antes del 05/10/2026 las semanas son lunes a domingo y desde el 05/10 son lunes a lunes (+7 días)."""
        # Septiembre 2026: todas deben terminar en domingo
        sep_weeks = self.manager._build_weeks(2026, 9)
        self.assertEqual(sep_weeks[0], (date(2026, 8, 31), date(2026, 9, 6)))
        self.assertEqual(sep_weeks[-1], (date(2026, 9, 28), date(2026, 10, 4)))
        for s, e in sep_weeks:
            self.assertEqual((e - s).days, 6, "Semanas de septiembre deben tener 7 días (lun a dom, delta=6)")

        # Octubre 2026: la semana puente (28-Sep a 04-Oct) termina en domingo; desde el 05/10 terminan en lunes
        oct_weeks = self.manager._build_weeks(2026, 10)
        self.assertEqual(oct_weeks[0], (date(2026, 9, 28), date(2026, 10, 4)), "Semana puente conserva término en domingo 04/10")
        self.assertEqual(oct_weeks[1], (date(2026, 10, 5), date(2026, 10, 12)))
        self.assertEqual(oct_weeks[2], (date(2026, 10, 12), date(2026, 10, 19)))
        self.assertEqual(oct_weeks[3], (date(2026, 10, 19), date(2026, 10, 26)))
        self.assertEqual(oct_weeks[4], (date(2026, 10, 26), date(2026, 11, 2)))

        # Verificar solapamiento: fin de semana N es inicio de semana N+1
        for i in range(1, len(oct_weeks) - 1):
            self.assertEqual(oct_weeks[i][1], oct_weeks[i + 1][0], "El lunes de fin debe ser el lunes de inicio del siguiente")

    def test_october_shifts_chaining_and_bridge_preservation(self):
        """Verifica que la semana puente de Octubre conserve la asignación de Septiembre y los nuevos turnos inicien el 05/10."""
        shifts = self.controller.preview_shifts(2026, 10, [])
        # Semana 0 (puente): conservada de septiembre (ROJAS PEDRO)
        self.assertEqual(shifts[0]["semana"], (date(2026, 9, 28), date(2026, 10, 4)))
        self.assertEqual(shifts[0]["persona"], "CBO (M) ROJAS PEDRO")

        # Semana 1 (05 al 12): SGT PEREZ JUAN (siguiente_id = 1)
        self.assertEqual(shifts[1]["semana"], (date(2026, 10, 5), date(2026, 10, 12)))
        self.assertEqual(shifts[1]["persona"], "SGT (F) PEREZ JUAN")

        # Semana 2 (12 al 19): CBO GOMEZ ANA
        self.assertEqual(shifts[2]["semana"], (date(2026, 10, 12), date(2026, 10, 19)))
        self.assertEqual(shifts[2]["persona"], "CBO (M) GOMEZ ANA")

    def test_handover_monday_exception_blocks_both_shifts(self):
        """
        Si un funcionario tiene una excepción el lunes de solapamiento (ej: 12/10),
        queda inhabilitado tanto para el turno que finaliza ese día (05/10 al 12/10)
        como para el turno que inicia ese día (12/10 al 19/10).
        """
        # PEREZ JUAN (id 1) normalmente tendría el turno 05/10 al 12/10.
        # Si tiene licencia el lunes 12/10, NO debe recibir el turno 05/10 al 12/10.
        exc_perez = [{"persona": "SGT (F) PEREZ JUAN", "fecha": date(2026, 10, 12), "tipo": "LIC"}]
        shifts = self.controller.preview_shifts(2026, 10, exc_perez)

        shift_5_to_12 = next(s for s in shifts if s["semana"][0] == date(2026, 10, 5))
        self.assertNotEqual(shift_5_to_12["persona"], "SGT (F) PEREZ JUAN")

        # Igualmente, si GOMEZ ANA (id 2) tiene excepción el 12/10, no puede tomar el turno 12/10 al 19/10.
        exc_gomez = [{"persona": "CBO (M) GOMEZ ANA", "fecha": date(2026, 10, 12), "tipo": "FL"}]
        shifts_gomez = self.controller.preview_shifts(2026, 10, exc_gomez)

        shift_12_to_19 = next(s for s in shifts_gomez if s["semana"][0] == date(2026, 10, 12))
        self.assertNotEqual(shift_12_to_19["persona"], "CBO (M) GOMEZ ANA")

    def test_excel_export_colors_handover_day_for_both_persons(self):
        """
        En la planilla Excel generada para Octubre 2026, el lunes 12 debe estar
        marcado en la fila de quien entrega (turno 05 al 12) y de quien recibe (12 al 19).
        """
        shifts = self.controller.preview_shifts(2026, 10, [])
        output_excel = self.root_path / "test_october_overlap.xlsx"
        handler = ExcelHandler(str(output_excel), self.personal)
        handler.load_template()
        handler.write_shifts(shifts, [], year=2026, month=10)
        handler.save_report()

        wb = openpyxl.load_workbook(str(output_excel))
        sheet = wb["Turnos"]

        # Perez Juan (id 1) está en turno 05 al 12 -> día 12 (columna 13) debe tener relleno
        # Gomez Ana (id 2) está en turno 12 al 19 -> día 12 (columna 13) debe tener relleno
        col_day_12 = 13  # A=1, Día 1=Col 2, Día 12=Col 13
        perez_cell = sheet.cell(row=4, column=col_day_12)  # Fila 4 es Perez
        gomez_cell = sheet.cell(row=5, column=col_day_12)  # Fila 5 es Gomez

        self.assertIsNotNone(perez_cell.fill.fill_type, "Perez Juan debe tener marcado el lunes 12")
        self.assertIsNotNone(gomez_cell.fill.fill_type, "Gomez Ana debe tener marcado el lunes 12")

        # Verificar totales de turnos: cada turno cuenta en el mes donde inicia.
        # En Octubre 2026 inician 4 turnos: 05/10, 12/10, 19/10, 26/10.
        # La semana puente 28-Sep a 04-Oct inició en septiembre, por lo que NO suma en octubre.
        turnos_totales_octubre = 0
        for r in range(4, 4 + len(self.personal)):
            turnos_totales_octubre += sheet.cell(row=r, column=33).value or 0
        self.assertEqual(turnos_totales_octubre, 4, "Octubre debe contabilizar exactamente los 4 turnos que inician en el mes")


if __name__ == "__main__":
    unittest.main()
