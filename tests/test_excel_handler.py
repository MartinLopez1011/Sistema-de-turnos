import os
import tempfile
import unittest
from datetime import date
from pathlib import Path
from unittest.mock import patch

import openpyxl

from utils.excel_handler import ExcelHandler
from controllers.main_controller import MainController


class ExcelHandlerTests(unittest.TestCase):
    def setUp(self):
        self.temp_dir = tempfile.TemporaryDirectory()
        self.output_path = Path(self.temp_dir.name) / "test_turnos.xlsx"
        self.personal = [
            {"id": 1, "nombre": "COM ALARCON JOSE"},
            {"id": 2, "nombre": "PRO BRAVO PATRICIO"},
        ]
        self.handler = ExcelHandler(str(self.output_path), self.personal)

    def tearDown(self):
        self.temp_dir.cleanup()

    def test_load_template_creates_correct_grid(self):
        self.handler.load_template()
        self.assertIsNotNone(self.handler.sheet)

        # Verificar título y encabezados
        self.assertEqual(self.handler.sheet["A1"].value, "PLANIFICACION DE TURNOS")
        self.assertEqual(self.handler.sheet["A2"].value, "DIA")
        self.assertEqual(self.handler.sheet["A3"].value, "FUNCIONARIO")

        # Verificar días 1 a 31
        for day in range(1, 32):
            self.assertEqual(self.handler.sheet.cell(row=2, column=day + 1).value, day)

        # Verificar filas de personal
        self.assertEqual(self.handler.sheet.cell(row=4, column=1).value, "COM ALARCON JOSE")
        self.assertEqual(self.handler.sheet.cell(row=5, column=1).value, "PRO BRAVO PATRICIO")

    def test_write_shifts_and_exceptions_and_save(self):
        self.handler.load_template()

        shifts = [
            {
                "semana": (date(2026, 9, 7), date(2026, 9, 13)),
                "persona": "COM ALARCON JOSE",
                "saltados": [],
            }
        ]
        exceptions = [
            {
                "persona": "PRO BRAVO PATRICIO",
                "fecha": date(2026, 9, 15),
                "tipo": "DA",
            }
        ]

        self.handler.write_shifts(shifts, exceptions, year=2026, month=9)
        self.handler.save_report()

        self.assertTrue(self.output_path.exists())

        # Reabrir el archivo y verificar contenido
        wb = openpyxl.load_workbook(self.output_path)
        sheet = wb["Turnos"]

        # La celda del día 15 de PRO BRAVO PATRICIO (fila 5) debe tener 'DA'
        # El día 15 está en la columna 16 (columna 1 es Funcionario, col 2 es día 1, ..., col 16 es día 15)
        day_15_cell = sheet.cell(row=5, column=16)
        self.assertEqual(day_15_cell.value, "DA")

    def test_permission_error_is_captured_gracefully_in_controller(self):
        # Crear un controller temporal
        config_path = Path(self.temp_dir.name) / "config.json"
        config_path.write_text('{"personal": [], "inicio": {}, "historial": {}, "siguiente_id": 1, "pendientes": [], "snapshots": {}, "excepciones": {}}', encoding="utf-8")
        controller = MainController(self.temp_dir.name)

        # Mockear save_report para simular archivo bloqueado por Excel
        with patch.object(ExcelHandler, "save_report", side_effect=PermissionError("Archivo bloqueado")):
            ok, msg = controller.process_generation(2026, 9, [])
            self.assertFalse(ok)
            self.assertIn("está abierto en Microsoft Excel", msg)


if __name__ == "__main__":
    unittest.main()
