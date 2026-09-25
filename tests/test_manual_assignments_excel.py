import os
import tempfile
import unittest
from datetime import date
from pathlib import Path
import openpyxl

from controllers.main_controller import MainController
from utils.excel_handler import ExcelHandler


class TestManualAssignmentsExcel(unittest.TestCase):
    def setUp(self):
        self.temp_dir = tempfile.TemporaryDirectory()
        self.root_path = Path(self.temp_dir.name)
        self.output_excel = str(self.root_path / "test_manual_turnos.xlsx")

        self.personal = [
            {"id": 1, "nombre": "Sargento Perez"},
            {"id": 2, "nombre": "Cabo Gomez"},
            {"id": 3, "nombre": "Soldado Lopez"},
            {"id": 4, "nombre": "Cabo Diaz"},
        ]

    def tearDown(self):
        self.temp_dir.cleanup()

    def test_manual_shift_cell_highlight_and_comment(self):
        """Verifica que el turno manual se pinte de verde esmeralda y tenga comentario con motivo."""
        shifts = [
            {
                "semana": (date(2026, 9, 7), date(2026, 9, 13)),
                "persona": "Cabo Gomez",
                "es_manual": True
            },
            {
                "semana": (date(2026, 9, 14), date(2026, 9, 20)),
                "persona": "Sargento Perez"
            }
        ]
        manual_changes = [
            {
                "semana": (date(2026, 9, 7), date(2026, 9, 13)),
                "nuevo": "Cabo Gomez",
                "anterior": "Sargento Perez",
                "motivo": "Accidente laboral de camino al cuartel"
            }
        ]

        with ExcelHandler(self.output_excel, self.personal) as handler:
            handler.load_template()
            handler.write_shifts(shifts, [], 2026, 9, manual_changes=manual_changes)
            handler.save_report()

        wb = openpyxl.load_workbook(self.output_excel)
        sheet = wb["Turnos"]

        # Encontrar fila de Cabo Gomez y fila de Sargento Perez en la grilla principal
        gomez_row = None
        perez_row = None
        for r in range(4, 4 + len(self.personal)):
            val = str(sheet.cell(row=r, column=1).value or "")
            if "Cabo Gomez" in val and gomez_row is None:
                gomez_row = r
            elif "Sargento Perez" in val and perez_row is None:
                perez_row = r

        self.assertIsNotNone(gomez_row)
        self.assertIsNotNone(perez_row)

        # Dia 8 de sept: Columna correspondiente al dia 8
        # En la fila de Cabo Gomez (cambio manual), debe ser verde esmeralda 059669
        # y tener comentario detallando asignado, original y motivo.
        cell_manual = None
        for c in range(2, 33):
            header_val = sheet.cell(row=2, column=c).value
            if header_val in (8, "8", 8.0, "08"):
                cell_manual = sheet.cell(row=gomez_row, column=c)
                break

        self.assertIsNotNone(cell_manual)
        self.assertIsNotNone(cell_manual.fill)
        self.assertIn("059669", str(cell_manual.fill.start_color.rgb).upper())
        self.assertIsNotNone(cell_manual.comment)
        comment_text = cell_manual.comment.text
        self.assertIn("CAMBIO MANUAL DE GUARDIA", comment_text)
        self.assertIn("Asignado: Cabo Gomez", comment_text)
        self.assertIn("Guardia original: Sargento Perez", comment_text)
        self.assertIn("Motivo: Accidente laboral de camino al cuartel", comment_text)

        # Dia 15 de sept: Columna correspondiente al dia 15
        # En la fila de Sargento Perez (turno regular), debe ser rojo FF3B30 y sin comentario de cambio manual
        cell_regular = None
        for c in range(2, 33):
            header_val = sheet.cell(row=2, column=c).value
            if header_val in (15, "15", 15.0):
                cell_regular = sheet.cell(row=perez_row, column=c)
                break

        self.assertIsNotNone(cell_regular)
        self.assertIn("FF3B30", str(cell_regular.fill.start_color.rgb).upper())
        self.assertIsNone(cell_regular.comment)

    def test_manual_changes_table_rendered_in_excel(self):
        """Verifica que la tabla de Registro de Cambios Manuales se dibuje correctamente al fondo de la hoja."""
        shifts = [
            {
                "semana": (date(2026, 9, 1), date(2026, 9, 6)),
                "persona": "Soldado Lopez",
                "es_manual": True
            }
        ]
        manual_changes = [
            {
                "semana": (date(2026, 9, 1), date(2026, 9, 6)),
                "nuevo": "Soldado Lopez",
                "anterior": "Cabo Diaz",
                "motivo": "Permuta acordada por emergencia familiar"
            }
        ]

        with ExcelHandler(self.output_excel, self.personal) as handler:
            handler.load_template()
            handler.write_shifts(shifts, [], 2026, 9, manual_changes=manual_changes)
            handler.save_report()

        wb = openpyxl.load_workbook(self.output_excel)
        sheet = wb["Turnos"]

        # Buscar título de la tabla
        title_found = False
        table_row = None
        for r in range(1, sheet.max_row + 1):
            val = str(sheet.cell(row=r, column=1).value or "")
            if "REGISTRO DE CAMBIOS MANUALES DE GUARDIA" in val:
                title_found = True
                table_row = r
                break

        self.assertTrue(title_found, "Debe existir el título 'REGISTRO DE CAMBIOS MANUALES DE GUARDIA (PERMUTAS / ACCIDENTES)'")

        # Fila encabezados (table_row + 1)
        hdr_row = table_row + 1
        self.assertEqual(sheet.cell(row=hdr_row, column=1).value, "FUNCIONARIO ASIGNADO")
        self.assertEqual(sheet.cell(row=hdr_row, column=2).value, "SEMANA")
        self.assertEqual(sheet.cell(row=hdr_row, column=6).value, "GUARDIA ORIGINAL")
        self.assertEqual(sheet.cell(row=hdr_row, column=12).value, "MOTIVO DEL CAMBIO / JUSTIFICACIÓN")

        # Fila datos (table_row + 2)
        data_row = table_row + 2
        self.assertEqual(sheet.cell(row=data_row, column=1).value, "Soldado Lopez")
        self.assertEqual(sheet.cell(row=data_row, column=2).value, "01/09 al 06/09")
        self.assertEqual(sheet.cell(row=data_row, column=6).value, "Cabo Diaz")
        self.assertEqual(sheet.cell(row=data_row, column=12).value, "Permuta acordada por emergencia familiar")

    def test_both_manual_changes_and_otr_tables_coexist(self):
        """Verifica que ambas tablas (Cambios Manuales y Excepciones OTR) se muestren ordenadas y sin solaparse."""
        shifts = [
            {
                "semana": (date(2026, 9, 1), date(2026, 9, 6)),
                "persona": "Cabo Gomez",
                "es_manual": True
            }
        ]
        exceptions = [
            {
                "persona": "Sargento Perez",
                "fecha": date(2026, 9, 10),
                "tipo": "OTR",
                "motivo": "Comisión de servicio especial"
            }
        ]
        manual_changes = [
            {
                "semana": (date(2026, 9, 1), date(2026, 9, 6)),
                "nuevo": "Cabo Gomez",
                "anterior": "Sargento Perez",
                "motivo": "Fractura por accidente en trayecto"
            }
        ]

        with ExcelHandler(self.output_excel, self.personal) as handler:
            handler.load_template()
            handler.write_shifts(shifts, exceptions, 2026, 9, manual_changes=manual_changes)
            handler.save_report()

        wb = openpyxl.load_workbook(self.output_excel)
        sheet = wb["Turnos"]

        mc_title_row = None
        otr_title_row = None

        for r in range(1, sheet.max_row + 1):
            val = str(sheet.cell(row=r, column=1).value or "")
            if "REGISTRO DE CAMBIOS MANUALES DE GUARDIA" in val:
                mc_title_row = r
            elif "DETALLE DE PERMISOS Y EXCEPCIONES ESPECIALES (OTR)" in val:
                otr_title_row = r

        self.assertIsNotNone(mc_title_row, "Tabla de cambios manuales debe existir")
        self.assertIsNotNone(otr_title_row, "Tabla OTR debe existir")
        # La tabla OTR debe estar ubicada debajo de la tabla de cambios manuales
        self.assertGreater(otr_title_row, mc_title_row)

    def test_controller_process_generation_builds_manual_changes(self):
        """Verifica que MainController.process_generation detecta las asignaciones manuales, calcula los reemplazos y genera el Excel con la tabla."""
        import json
        config_path = self.root_path / "config.json"
        config_data = {
            "personal": self.personal,
            "inicio": {},
            "historial": {},
            "siguiente_id": 1,
            "pendientes": [],
            "snapshots": {},
            "excepciones": {},
            "asignaciones_manuales": {},
            "asignaciones_manuales_motivos": {}
        }
        config_path.write_text(json.dumps(config_data), encoding="utf-8")

        controller = MainController(str(self.root_path))
        # Previsualizar para obtener la clave de la semana 1
        shifts_preview = controller.preview_shifts(2026, 9, [])
        s_d, e_d = shifts_preview[0]["semana"]
        week_key = f"{s_d.isoformat()}_{e_d.isoformat()}"
        original_person = shifts_preview[0]["persona"]

        # Forzar a otra persona en esa semana con motivo
        new_person = "Soldado Lopez" if original_person != "Soldado Lopez" else "Cabo Diaz"
        manual_dict = {week_key: new_person}
        motives_dict = {week_key: "Sustitución por esguince"}

        ok, msg = controller.process_generation(
            2026, 9, [],
            manual_assignments=manual_dict,
            manual_motives=motives_dict,
            target_path=self.output_excel
        )
        self.assertTrue(ok, msg)

        # Cargar Excel y verificar
        wb = openpyxl.load_workbook(self.output_excel)
        sheet = wb["Turnos"]

        # Verificar que la tabla de cambios manuales se creó
        found_change_in_sheet = False
        for r in range(1, sheet.max_row + 1):
            if sheet.cell(row=r, column=1).value == new_person:
                # Comprobar si en las columnas siguientes figura el original y el motivo
                orig_val = sheet.cell(row=r, column=6).value
                mot_val = sheet.cell(row=r, column=12).value
                if orig_val == original_person and mot_val == "Sustitución por esguince":
                    found_change_in_sheet = True
                    break

        self.assertTrue(found_change_in_sheet, "El cambio manual debe estar registrado en la tabla de Excel con el funcionario anterior y el motivo")


if __name__ == "__main__":
    unittest.main()
