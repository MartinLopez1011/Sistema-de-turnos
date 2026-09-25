import os
import json
import tempfile
import unittest
from datetime import date
from pathlib import Path
from unittest.mock import MagicMock, patch

import openpyxl

from models.shift_manager import ShiftManager
from controllers.main_controller import MainController
from utils.excel_handler import ExcelHandler
from views.components.widgets import _make_row_hover
from main import setup_global_exception_handler


class ImprovementsTests(unittest.TestCase):
    def setUp(self):
        self.temp_dir = tempfile.TemporaryDirectory()
        self.root_path = Path(self.temp_dir.name)
        self.config_path = self.root_path / "config.json"
        self.output_excel = self.root_path / "test_excel_improvements.xlsx"

        self.personal = [
            {"id": 1, "nombre": "COM PEREZ JUAN"},
            {"id": 2, "nombre": "SGT GOMEZ ANA"},
        ]
        self.data = {
            "personal": self.personal,
            "inicio": {},
            "historial": {},
            "siguiente_id": 1,
            "pendientes": [],
            "snapshots": {},
            "excepciones": {},
            "asignaciones_manuales": {}
        }
        self.config_path.write_text(json.dumps(self.data), encoding="utf-8")
        self.manager = ShiftManager(str(self.config_path))
        self.controller = MainController(str(self.root_path))

    def tearDown(self):
        self.temp_dir.cleanup()

    def test_excel_weekdays_and_page_setup(self):
        handler = ExcelHandler(str(self.output_excel), self.personal)
        handler.load_template()

        # Verificar orientación apaisada y fit-to-page
        self.assertEqual(handler.sheet.page_setup.orientation, "landscape")
        self.assertEqual(handler.sheet.page_setup.fitToWidth, 1)
        self.assertEqual(handler.sheet.page_setup.fitToHeight, 0)
        self.assertTrue(handler.sheet.sheet_properties.pageSetUpPr.fitToPage)

        # Escribir turnos para Septiembre 2026 (1 de Septiembre de 2026 fue Martes -> "M")
        handler.write_shifts([], [], year=2026, month=9)
        handler.save_report()

        wb = openpyxl.load_workbook(str(self.output_excel))
        sheet = wb["Turnos"]

        # Fila 3, Columna 2 es el día 1 de Septiembre 2026 (Martes -> "M")
        day_1_letter = sheet.cell(row=3, column=2).value
        self.assertEqual(day_1_letter, "M")

        # Fila 3, Columna 3 es el día 2 de Septiembre 2026 (Miércoles -> "X")
        day_2_letter = sheet.cell(row=3, column=3).value
        self.assertEqual(day_2_letter, "X")

        # Fila 3, Columna 32 es el día 31 de Septiembre 2026 (no existe -> "-")
        day_31_letter = sheet.cell(row=3, column=32).value
        self.assertEqual(day_31_letter, "-")

    def test_excel_excludes_funcionario_header_from_person_rows(self):
        handler = ExcelHandler(str(self.output_excel), self.personal)
        handler.load_template()
        handler.write_shifts([], [], year=2026, month=9)

        # Verificar que A3 conserva 'FUNCIONARIO'
        self.assertEqual(handler.sheet.cell(row=3, column=1).value, "FUNCIONARIO")

    def test_exception_date_string_normalization(self):
        # Pasar excepciones con fecha como string ("2026-09-01") en vez de date
        exceptions = [
            {"persona": "COM PEREZ JUAN", "fecha": "2026-09-01", "tipo": "DA"}
        ]
        # No debe lanzar TypeError y debe procesar la excepción
        shifts, final_id, pendientes = self.manager.generate_shifts(2026, 9, exceptions)
        self.assertTrue(len(shifts) > 0)
        # La semana 1 (31 ago - 6 sep) le tocaba a Perez Juan, pero se salta por tener excepción el día 1
        week_1 = shifts[0]
        self.assertNotEqual(week_1['persona'], "COM PEREZ JUAN")
        self.assertEqual(week_1['persona'], "SGT GOMEZ ANA")
        self.assertTrue(any(s['persona'] == "COM PEREZ JUAN" for s in week_1['saltados']))

    def test_was_forced_recently_checks_manual_assignments(self):
        # Registrar una asignación manual en agosto 2026
        self.manager.asignaciones_manuales = {
            "2026-08": {
                "2026-08-10_2026-08-16": "COM PEREZ JUAN"
            }
        }
        # Evaluar si fue forzado/manual recientemente respecto al 24 de agosto de 2026 (dentro de 4 semanas)
        was_forced = self.manager._was_forced_recently("COM PEREZ JUAN", date(2026, 8, 24), [])
        self.assertTrue(was_forced)

        # Para otra persona debe dar False
        self.assertFalse(self.manager._was_forced_recently("SGT GOMEZ ANA", date(2026, 8, 24), []))

    def test_make_row_hover_supports_fg_color(self):
        # Simular widget CTkFrame (soporta fg_color) y tk.Frame (soporta bg)
        ctk_mock = MagicMock()
        ctk_mock.winfo_children.return_value = []
        ctk_mock.configure.side_effect = lambda **kwargs: None

        ref_mock = MagicMock()
        _make_row_hover(ref_mock, [(ctk_mock, "#1C2228")])
        # Verificar que se asociaron los eventos <Enter> y <Leave>
        self.assertTrue(ctk_mock.bind.called)

    def test_edit_person_updates_asignaciones_manuales(self):
        # Configurar asignación manual con el nombre actual
        self.manager.asignaciones_manuales = {
            "2026-09": {
                "2026-09-07_2026-09-13": "COM PEREZ JUAN"
            }
        }
        # Editar nombre de la persona 1
        success = self.manager.edit_person(1, "COM PEREZ JUAN CARLOS")
        self.assertTrue(success)
        # Verificar que la asignación manual se actualizó con el nuevo nombre
        self.assertEqual(
            self.manager.asignaciones_manuales["2026-09"]["2026-09-07_2026-09-13"],
            "COM PEREZ JUAN CARLOS"
        )

    def test_load_config_non_dict_failsafe(self):
        # Guardar un JSON que es una lista en vez de un diccionario
        self.config_path.write_text(json.dumps(["item1", "item2"]), encoding="utf-8")
        # Al recargar, debe recuperarse de forma segura y usar defaults limpios
        self.manager.load_config()
        self.assertEqual(self.manager.personal, [])
        self.assertEqual(self.manager.siguiente_id, 1)

    def test_excel_invalid_days_font_contrast(self):
        handler = ExcelHandler(str(self.output_excel), self.personal)
        handler.load_template()
        # Septiembre 2026 no tiene día 31 (debe tener fill 595959 y texto blanco FFFFFF)
        handler.write_shifts([], [], year=2026, month=9)
        handler.save_report()

        wb = openpyxl.load_workbook(str(self.output_excel))
        sheet = wb["Turnos"]
        cell_31 = sheet.cell(row=2, column=32)
        self.assertEqual(cell_31.value, "-")
        self.assertIn("FFFFFF", str(cell_31.font.color.rgb))
        self.assertTrue(cell_31.font.bold)

    def test_tab_settings_empty_personal_dropdown(self):
        from views.tabs.tab_settings import TabSettings
        mock_parent = MagicMock()
        mock_app = MagicMock()
        mock_app.controller = MagicMock()
        mock_app.controller.get_personal_list.return_value = []
        mock_app.controller.get_all_persons.return_value = []

        tab_settings = TabSettings.__new__(TabSettings)
        tab_settings.app = mock_app
        tab_settings.controller = mock_app.controller
        tab_settings.starting_person_var = MagicMock()
        tab_settings.starting_person_dropdown = MagicMock()
        tab_settings.person_list_frame = MagicMock()
        tab_settings.person_list_frame.winfo_children.return_value = []

        tab_settings.refresh_person_list()
        tab_settings.starting_person_dropdown.configure.assert_called_with(values=["Sin personal disponible"])
        tab_settings.starting_person_var.set.assert_called_with("Sin personal disponible")

    def test_global_exception_handler_captures_exceptions(self):
        import sys
        old_hook = sys.excepthook
        try:
            setup_global_exception_handler()
            with patch("main.logger.critical") as mock_log, patch("main.messagebox.showerror") as mock_box:
                try:
                    raise ValueError("Prueba de error no controlado")
                except ValueError as e:
                    sys.excepthook(type(e), e, e.__traceback__)
                    mock_log.assert_called_once()
                    mock_box.assert_called_once()
        finally:
            sys.excepthook = old_hook

    @patch("views.tabs.tab_settings.ctk.CTkFont")
    @patch("views.tabs.tab_settings._make_row_hover")
    @patch("views.tabs.tab_settings.ctk.CTkButton")
    @patch("views.tabs.tab_settings.ctk.CTkLabel")
    @patch("views.tabs.tab_settings.ctk.CTkFrame")
    def test_tab_settings_person_count_badge(self, mock_frame, mock_label, mock_btn, mock_hover, mock_font):
        from views.tabs.tab_settings import TabSettings
        mock_app = MagicMock()
        mock_app.controller = MagicMock()
        mock_app.controller.get_personal_list.return_value = ["P1", "P2"]
        mock_app.controller.get_all_persons.return_value = [{"id": 1, "nombre": "P1"}, {"id": 2, "nombre": "P2"}]

        tab_settings = TabSettings.__new__(TabSettings)
        tab_settings.app = mock_app
        tab_settings.controller = mock_app.controller
        tab_settings.starting_person_var = MagicMock()
        tab_settings.starting_person_dropdown = MagicMock()
        tab_settings.person_list_frame = MagicMock()
        tab_settings.person_list_frame.winfo_children.return_value = []
        tab_settings.person_count_badge = MagicMock()

        tab_settings.refresh_person_list()
        tab_settings.person_count_badge.configure.assert_called_once()
        badge_text = tab_settings.person_count_badge.configure.call_args[1]["text"]
        self.assertIn("2 funcionarios activos", badge_text)

    @patch("views.tabs.tab_plan.ctk.CTkFont")
    @patch("views.tabs.tab_plan._avatar_ctk")
    @patch("views.tabs.tab_plan.ctk.CTkLabel")
    @patch("views.tabs.tab_plan.ctk.CTkFrame")
    def test_tab_plan_render_next_turno_card(self, mock_frame, mock_label, mock_avatar, mock_font):
        from views.tabs.tab_plan import TabPlan
        mock_app = MagicMock()
        tab_plan = TabPlan.__new__(TabPlan)
        tab_plan.app = mock_app
        tab_plan.next_turno_frame = MagicMock()
        tab_plan.next_turno_frame.winfo_children.return_value = []

        shift = {
            "semana": (date(2026, 9, 7), date(2026, 9, 13)),
            "persona": "COM PEREZ JUAN",
        }
        tab_plan._render_next_turno_card(shift, date(2026, 9, 8), ["COM PEREZ JUAN"])
        self.assertIsNotNone(tab_plan.next_turno_lbl)
        # Verify text was set to PEREZ JUAN
        label_texts = [call[1].get("text") for call in mock_label.call_args_list if "text" in call[1]]
        self.assertIn("PEREZ JUAN", label_texts)

    @patch("views.tabs.tab_plan.ctk.CTkFont")
    @patch("views.tabs.tab_plan.ctk.CTkLabel")
    @patch("views.tabs.tab_plan.ctk.CTkFrame")
    def test_tab_plan_update_preview_empty_shifts_no_unbound_local_error(self, mock_frame, mock_label, mock_font):
        from views.tabs.tab_plan import TabPlan
        mock_app = MagicMock()
        mock_app.controller = MagicMock()
        mock_app.controller.get_personal_list.return_value = ["COM PEREZ JUAN"]

        tab_plan = TabPlan.__new__(TabPlan)
        tab_plan.app = mock_app
        tab_plan.controller = mock_app.controller
        tab_plan.preview_scroll = MagicMock()
        tab_plan.preview_scroll.winfo_children.return_value = []
        tab_plan._render_next_turno_card = MagicMock()

        # Debe ejecutarse sin lanzar UnboundLocalError
        tab_plan.update_preview([], 2026, 9)
        tab_plan._render_next_turno_card.assert_called_once()

    def test_parse_week_range_cache_and_correctness(self):
        from models.shift_manager import _parse_week_range
        # Clave válida
        s, e = _parse_week_range("2026-09-07_2026-09-13")
        self.assertEqual(s, date(2026, 9, 7))
        self.assertEqual(e, date(2026, 9, 13))

        # Clave inválida
        s_inv, e_inv = _parse_week_range("invalido_formato")
        self.assertIsNone(s_inv)
        self.assertIsNone(e_inv)

        # Verificar que el caché LRU registra hits
        info_before = _parse_week_range.cache_info()
        _parse_week_range("2026-09-07_2026-09-13")
        info_after = _parse_week_range.cache_info()
        self.assertGreater(info_after.hits, info_before.hits)

    def test_excel_summary_totals_and_dynamic_column_width(self):
        # Personal con nombre largo para verificar ancho dinámico
        long_name_personal = [
            {"id": 1, "nombre": "COORDINADOR GENERAL DEPARTAMENTO ADMINISTRATIVO PEREZ GONZALEZ JUAN CARLOS"},
            {"id": 2, "nombre": "GOMEZ ANA"},
        ]
        handler = ExcelHandler(str(self.output_excel), long_name_personal)
        handler.load_template()

        # Verificar auto-ancho de columna A
        expected_width = len("COORDINADOR GENERAL DEPARTAMENTO ADMINISTRATIVO PEREZ GONZALEZ JUAN CARLOS") + 4
        self.assertEqual(handler.sheet.column_dimensions["A"].width, expected_width)

        # Verificar anchos de columnas de resumen (AG a AK)
        self.assertEqual(handler.sheet.column_dimensions["AG"].width, 10)
        self.assertEqual(handler.sheet.column_dimensions["AH"].width, 6)

        # Verificar encabezados de resumen en filas 2 y 3
        self.assertEqual(handler.sheet.cell(row=2, column=33).value, "TOTALES")
        self.assertEqual(handler.sheet.cell(row=3, column=33).value, "TURNOS")
        self.assertEqual(handler.sheet.cell(row=3, column=34).value, "DA")
        self.assertEqual(handler.sheet.cell(row=3, column=35).value, "FL")
        self.assertEqual(handler.sheet.cell(row=3, column=36).value, "LIC")
        self.assertEqual(handler.sheet.cell(row=3, column=37).value, "OTR")

        # Escribir turnos y excepciones
        shifts = [
            {
                "semana": (date(2026, 9, 7), date(2026, 9, 13)),
                "persona": long_name_personal[0]["nombre"],
            }
        ]
        exceptions = [
            {"persona": long_name_personal[0]["nombre"], "fecha": date(2026, 9, 1), "tipo": "DA"},
            {"persona": long_name_personal[0]["nombre"], "fecha": date(2026, 9, 2), "tipo": "DA"},
            {"persona": long_name_personal[1]["nombre"], "fecha": date(2026, 9, 15), "tipo": "FL"},
            {"persona": long_name_personal[1]["nombre"], "fecha": date(2026, 9, 16), "tipo": "LIC"},
        ]

        handler.write_shifts(shifts, exceptions, year=2026, month=9)
        handler.save_report()

        wb = openpyxl.load_workbook(str(self.output_excel))
        sheet = wb["Turnos"]

        # Fila 4: long_name_personal[0] -> 1 turno, 2 DA, 0 FL, 0 LIC, 0 OTR
        self.assertEqual(sheet.cell(row=4, column=33).value, 1)  # TURNOS
        self.assertEqual(sheet.cell(row=4, column=34).value, 2)  # DA
        self.assertEqual(sheet.cell(row=4, column=35).value, 0)  # FL
        self.assertEqual(sheet.cell(row=4, column=36).value, 0)  # LIC
        self.assertEqual(sheet.cell(row=4, column=37).value, 0)  # OTR

        # Fila 5: long_name_personal[1] -> 0 turnos, 0 DA, 1 FL, 1 LIC, 0 OTR
        self.assertEqual(sheet.cell(row=5, column=33).value, 0)  # TURNOS
        self.assertEqual(sheet.cell(row=5, column=34).value, 0)  # DA
        self.assertEqual(sheet.cell(row=5, column=35).value, 1)  # FL
        self.assertEqual(sheet.cell(row=5, column=36).value, 1)  # LIC
        self.assertEqual(sheet.cell(row=5, column=37).value, 0)  # OTR

    def test_tab_settings_github_actions(self):
        from views.tabs.tab_settings import TabSettings
        from unittest.mock import patch, MagicMock

        tab_settings = TabSettings.__new__(TabSettings)
        mock_parent = MagicMock()
        mock_app = MagicMock()
        mock_btn = MagicMock()

        tab_settings.parent = mock_parent
        tab_settings.app = mock_app
        tab_settings.btn_copy_github = mock_btn

        with patch("webbrowser.open_new_tab") as mock_open:
            tab_settings._open_github()
            mock_open.assert_called_once_with("https://github.com/MartinLopez1011/Sistema-de-turnos.git")

        tab_settings._copy_github_link()
        mock_parent.clipboard_clear.assert_called_once()
        mock_parent.clipboard_append.assert_called_once_with("https://github.com/MartinLopez1011/Sistema-de-turnos.git")
        mock_btn.configure.assert_called_once()


if __name__ == "__main__":
    unittest.main()






